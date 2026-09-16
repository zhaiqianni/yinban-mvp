from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
ASSET_DIR = BASE_DIR / "assets"
OUTPUT_PATH = BASE_DIR / "银伴医院智能陪诊导引机器人项目计划书.docx"

FONT_CN = "Microsoft YaHei"
FONT_EN = "Aptos"
BLACK = "000000"
DEEP_TEAL = "075448"
DARK_HEADER = "24413B"
GOLD = "F2AD32"
PALE = "F4F6F5"
PALE_GOLD = "FFF7E2"
WHITE = "FFFFFF"
BORDER = "D9D9D9"
MUTED = "5F6F6A"


def set_run_font(run, size=None, bold=None, color=BLACK, name=FONT_CN):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_EN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_picture_with_alt(run, path, *, alt, width=None, height=None):
    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    picture = run.add_picture(str(path), **kwargs)
    picture._inline.docPr.set("descr", alt)
    picture._inline.docPr.set("title", alt)
    return picture


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    repeat_header(row)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, *, header_fill=DARK_HEADER, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, label in enumerate(headers):
        set_cell_text(
            header.cells[idx],
            label,
            bold=True,
            color=WHITE,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(header.cells[idx], header_fill)
        if widths:
            header.cells[idx].width = widths[idx]
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_idx, value in enumerate(values):
            center = col_idx == 0 and len(headers) > 2
            if isinstance(value, tuple):
                text, center = value
            else:
                text = value
            set_cell_text(
                row.cells[col_idx],
                text,
                size=font_size,
                align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT,
            )
            if widths:
                row.cells[col_idx].width = widths[col_idx]
            if row_idx % 2 == 1:
                set_cell_shading(row.cells[col_idx], PALE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    text = re.sub(r"^(\d+)\s+(\d+)\s+(\d+)(?=\s)", r"\1.\2.\3", text)
    text = re.sub(r"^(\d+)\s+(\d+)(?=\s)", r"\1.\2", text)
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    run = p.add_run(text)
    heading_size = {1: 18, 2: 13, 3: 11}.get(level, 10.8)
    set_run_font(run, size=heading_size, bold=True, color=BLACK)
    return p


def add_para(doc, text="", *, bold_lead=None, align=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    # 正文按中文项目计划书惯例统一首行缩进两个字符。
    p.paragraph_format.first_line_indent = Pt(21.6)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.28
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=10.8, bold=True)
        text = text[len(bold_lead):]
    run = p.add_run(text)
    set_run_font(run, size=10.8, color=BLACK)
    run.italic = italic
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        set_run_font(p.add_run(item), size=10.5)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        set_run_font(p.add_run(f"{index}.  {item}"), size=10.5)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=MUTED)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), DEEP_TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_EN)
    r_fonts.set(qn("w:hAnsi"), FONT_EN)
    r_fonts.set(qn("w:eastAsia"), FONT_CN)
    r_pr.append(r_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("银伴项目计划书  "), size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instr)
    run.append(separate)
    run.append(value)
    run.append(end)


def add_section_break(doc):
    # Let Word paginate chapters naturally so short closing subsections do not
    # strand a mostly empty page. Heading styles keep titles with their content.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    title = doc.styles["Title"]
    title.font.name = FONT_CN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_p_pr = title._element.get_or_add_pPr()
    for border in title_p_pr.findall(qn("w:pBdr")):
        title_p_pr.remove(border)

    for name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    add_picture_with_alt(
        p.add_run(),
        ASSET_DIR / "yinban-logo-horizontal-transparent.png",
        width=Cm(10.5),
        alt="银伴横向Logo：友好机器人、银伴中文名与YINBAN英文名",
    )

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run("银伴医院智能陪诊导引机器人项目计划书"), size=26, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("创新创业比赛通用版"), size=13, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_picture_with_alt(
        p.add_run(),
        ASSET_DIR / "yinban-mascot-cover-transparent.png",
        height=Cm(10.2),
        alt="银伴机器人封面形象：奶白机器人、墨绿面屏和暖金导航环",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run("让复杂就医变得有人陪伴"), size=14, bold=True, color=DEEP_TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("项目阶段  医院固定环境参赛原型"), size=10, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("2026年9月"), size=10, color=MUTED)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "目录", 1)
    entries = [
        ("1", "项目摘要"),
        ("2", "需求背景"),
        ("3", "用户与产品方案"),
        ("4", "原型成果与技术架构"),
        ("5", "硬件系统与安全设计"),
        ("6", "创新与竞争分析"),
        ("7", "市场与商业路径"),
        ("8", "实施与资源计划"),
        ("9", "验证方案与风险管理"),
        ("10", "社会价值与发展边界"),
        ("附录", "实施清单与资料来源"),
    ]
    for chapter, title in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.18
        set_run_font(p.add_run(f"{chapter}  {title}"), size=10.8)

    add_heading(doc, "阅读说明", 2)
    add_para(
        doc,
        "本计划书面向创新创业比赛评审。文中把已经通过代码与自动测试验证的功能、硬件到货后需要完成的现场验证，以及未来产品规划分开表述。市场与预算中的预测均注明计算假设，不作为已经实现的经营结果。",
    )
    add_para(
        doc,
        "项目当前的核心判断是：在比赛周期内，团队应证明老年友好交互、医院流程问答、路线展示和实体巡线小车能够形成稳定闭环；真实医院部署所需的通用定位、动态导航、电梯联动和医院信息系统接口属于后续阶段。",
        bold_lead="项目当前的核心判断是：",
    )


def build_document():
    doc = Document()
    configure_document(doc)
    cp = doc.core_properties
    cp.title = "银伴医院智能陪诊导引机器人项目计划书"
    cp.subject = "创新创业比赛项目计划书"
    cp.author = "银伴项目"
    cp.keywords = "适老化 医院导引 机器人 ESP32 FastAPI"

    add_cover(doc)
    add_contents(doc)

    doc.add_page_break()
    add_heading(doc, "1 项目摘要", 1)
    add_heading(doc, "1 1 项目结论", 2)
    add_para(
        doc,
        "银伴是一套面向老年就医者的医院智能陪诊导引原型。用户可以直接说出目的地或点击大按钮，系统给出科室位置、就医步骤和可视路线；在支持的实体路线中，ESP32巡线小车接收任务并完成低速导引、障碍停车和到达反馈。项目把问询、流程辅助、路线展示与实体行动放在同一个入口中，重点解决复杂公共空间中的信息理解和下一步行动问题。",
    )
    add_para(
        doc,
        "目前软件端已经形成可独立演示的闭环，包含适老界面、语音与文字输入、五类地点问询、按说话顺序解析多目的地、真实道路组合路线、模拟机器人以及启动停止协议。2026年9月9日在项目目录运行自动测试，结果为26项通过。硬件已完成采购，下一阶段是到货核验、引脚映射、巡线标定和软硬件联调。",
        bold_lead="目前软件端已经形成可独立演示的闭环，",
    )
    add_para(
        doc,
        "本项目的参赛价值在于低成本、可解释和可验证。它不依赖在线大模型完成核心问询，不把固定线路包装成通用自主导航，也不接触真实患者身份或诊疗数据。比赛现场即使语音、网络或蓝牙出现故障，文字输入、按钮操作、USB串口和模拟机器人仍可维持主流程。",
    )

    add_heading(doc, "1 2 当前阶段证据", 2)
    add_table(
        doc,
        ["类别", "当前状态", "可核验依据"],
        [
            ("软件", "已完成", "网页交互、问答、路线规划、模拟机器人和API均已实现"),
            ("测试", "已完成", "pytest自动测试26项通过，覆盖API、意图、路线、协议和模拟状态"),
            ("品牌", "已完成", "原创银伴机器人图标、报名方形版、横向字标和封面形象"),
            ("采购", "已完成", "硬件进入到货核验和组装阶段，预算上限为500元"),
            ("实物联调", "待验证", "需要按实物型号确认引脚、电平、电机方向和传感器阈值"),
        ],
        widths=[Cm(2.2), Cm(2.4), Cm(12.1)],
    )

    add_heading(doc, "1 3 能力边界", 2)
    add_bullets(
        doc,
        [
            "当前原型使用预设医院地图和固定路线，不具备SLAM或陌生环境自主建图能力。",
            "当前实体路线只承诺一条主路线，其他目的地先在屏幕中模拟导引。",
            "系统提供位置与流程信息，不提供诊断、处方、药物剂量或治疗建议。",
            "人工求助目前只触发本地提醒，不宣称已经通知医院工作人员。",
            "原型不读取身份证、医保卡、病历、人脸或真实预约数据。",
        ],
    )

    add_section_break(doc)
    add_heading(doc, "2 需求背景", 1)
    add_heading(doc, "2 1 人口与公共服务背景", 2)
    p = add_para(
        doc,
        "国家统计局发布的2025年国民经济和社会发展统计公报显示，2025年末我国60岁及以上人口为32338万人，占总人口23.0%；其中65岁及以上人口为22365万人，占15.9%。同一公报记录全年医疗卫生机构总诊疗人次为105.8亿。人口结构和服务规模共同说明，面向老年人的医院流程支持具有长期而广泛的应用基础。",
    )
    set_run_font(p.add_run(" [1]"), size=8.5, color=DEEP_TEAL).font.superscript = True
    add_para(
        doc,
        "这些数字不能直接推导出机器人的市场规模，也不能说明所有老年人都存在数字困难。项目据此判断的是：医院属于高频、流程密集且容易产生信息负担的公共场景，值得用可用性测试验证新的辅助方式。",
    )

    add_heading(doc, "2 2 就医旅程中的具体问题", 2)
    add_table(
        doc,
        ["阶段", "常见困难", "银伴的对应方式"],
        [
            ("到院", "入口多、楼栋和楼层关系陌生", "用自然语言确认目的地并展示简化地图"),
            ("挂号报到", "自助机步骤多，窗口与机器分流不清", "解释常见流程并明确人工窗口兜底"),
            ("候诊检查", "下一站信息分散，容易遗漏顺序", "按说话顺序管理多个目的地"),
            ("缴费取药", "位置变化和凭证要求难记", "同时给出地点、方向和非诊疗提醒"),
            ("迷路不适", "焦虑时难以操作复杂界面", "保留大按钮停止和本地人工求助入口"),
        ],
        widths=[Cm(2.6), Cm(6.0), Cm(8.1)],
    )

    add_heading(doc, "2 3 政策依据", 2)
    p = add_para(
        doc,
        "国务院办公厅2020年发布的实施方案要求在涉及老年人的高频服务场景中保留传统服务，并推动解决老年人运用智能技术的突出困难。文件特别提出医疗机构应保留人工服务窗口并配备导医等人员。银伴的设计因此不是用机器人替代人工，而是让机器人承担重复问询和固定路线辅助，把人工资源留给复杂问题。",
    )
    set_run_font(p.add_run(" [2]"), size=8.5, color=DEEP_TEAL).font.superscript = True
    p = add_para(
        doc,
        "国家卫生健康委和国家中医药局发布的改善就医感受提升患者体验主题活动方案提出从患者视角梳理全过程服务，运用新技术打通流程中的堵点。银伴把适老界面、流程说明和导引行动连接起来，与这一方向一致。该文件实施期为2023至2025年，本计划书将其作为政策背景，而非2026年的新增政策。",
    )
    set_run_font(p.add_run(" [3]"), size=8.5, color=DEEP_TEAL).font.superscript = True
    p = add_para(
        doc,
        "中华人民共和国无障碍环境建设法要求医院等公共服务场所的自助服务终端具备语音、大字等无障碍功能，并在医疗卫生等场所保留现场指导和人工办理等传统服务方式。该要求支持银伴采用大字体、语音反馈和人工接管，但不意味着机器人可以替代医院现有服务责任。",
    )
    set_run_font(p.add_run(" [4]"), size=8.5, color=DEEP_TEAL).font.superscript = True
    p = add_para(
        doc,
        "国家卫生健康委2025年有关建议答复进一步提出优化老年人就医流程、推进智能设备适老化，并发挥志愿服务作用。银伴据此把产品定位为医院服务人员的辅助工具，同时保留人工窗口、志愿者和导诊人员的兜底渠道。",
    )
    set_run_font(p.add_run(" [5]"), size=8.5, color=DEEP_TEAL).font.superscript = True

    add_section_break(doc)
    add_heading(doc, "3 用户与产品方案", 1)
    add_heading(doc, "3 1 目标用户", 2)
    add_para(
        doc,
        "首批目标用户是在大型医院中需要方向和流程帮助的老年就医者，尤其是不熟悉智能手机和自助机、视力或听力下降、第一次到院或需要跨多个科室行动的人。陪同家属、志愿者和导诊人员也是间接用户，他们可以用银伴确认路线并减少重复问询。",
    )
    add_table(
        doc,
        ["用户", "任务", "设计要求"],
        [
            ("老年就医者", "问地点、问流程、跟随导引、停止或求助", "大字体、高对比、慢语速、少步骤、多入口"),
            ("陪同家属", "安排多个目的地并确认顺序", "一次输入多站点，路线与文字保持一致"),
            ("导诊与志愿者", "处理高频重复问题", "知识库可维护，异常时转人工"),
            ("医院管理者", "评估服务效率与风险", "可解释规则、匿名日志、明确能力边界"),
        ],
        widths=[Cm(3.0), Cm(6.2), Cm(7.5)],
    )

    add_heading(doc, "3 1 1 利益相关者与真实使用场景", 3)
    add_para(
        doc,
        "首轮调研不直接询问受访者是否喜欢机器人，而是围绕一次真实就医经历追问任务、困难、替代做法和拒绝原因。这样可以区分礼貌性认可与真实使用意愿，并发现项目当前未覆盖的障碍。",
    )
    add_table(
        doc,
        ["对象", "优先场景", "需要收集的证据"],
        [
            ("老年就医者", "首次到院 跨科室 检查后取药", "最困难步骤 当前解决办法 操作障碍 是否愿意跟随"),
            ("陪同家属", "多站点安排 临时改路线", "顺序管理需求 等待成本 对安全和隐私的顾虑"),
            ("导诊 志愿者 护士", "大厅问询 人工窗口 异常求助", "高频问题 禁用表述 人工接管触发条件"),
            ("医院管理者", "封闭试点 设备运维", "允许运行区域 地图责任 日志边界 验收标准"),
        ],
        widths=[Cm(3.2), Cm(5.1), Cm(8.4)],
        font_size=8.8,
    )

    add_heading(doc, "3 1 2 用户访谈计划与关键假设", 3)
    add_para(
        doc,
        "首轮计划访谈5至8名老年就医者或陪同家属；条件允许时，再访谈1至2名导诊、护士或志愿者。记录年龄段、就医频率、最困难任务、当前解决方式、跟随低速机器人的意愿与顾虑。所有记录匿名保存，同时保留拒绝意见和失败样本。",
    )
    add_table(
        doc,
        ["待验证假设", "测试方法", "首轮判断标准"],
        [
            ("大按钮比多层菜单更易完成任务", "同一任务分别使用两种界面，记录完成时间、误触和求助次数", "多数受试者用大按钮更快完成且误触不增加"),
            ("语音入口能降低输入负担", "在安静与模拟大厅噪声下说出单站和多站请求", "识别文本可确认；失败时能立即改用按钮或文字"),
            ("实体低速导引比静态地图更容易执行", "比较看图自行前往与跟随小车两种方式", "跟随组少走错路，且受试者未出现明显不安"),
            ("目标用户愿意在受控场景跟随机器人", "展示启动、暂停、避障和人工接管后询问接受条件", "记录接受、拒绝及原因，不把口头好评等同使用意愿"),
        ],
        widths=[Cm(4.6), Cm(7.1), Cm(5.0)],
        font_size=8.5,
    )

    add_heading(doc, "3 2 产品定位", 2)
    add_para(
        doc,
        "银伴的产品定义是面向复杂公共空间的适老陪伴导引终端。医院版本先完成五类能力：问地点、办流程、带路线、陪等待、接求助。当前原型重点验证前三类能力，并为后两类保留清晰而安全的入口。",
    )
    add_table(
        doc,
        ["问", "办", "带", "陪", "助"],
        [
            (
                "科室与设施位置",
                "挂号到取药流程",
                "地图与实体导引",
                "候诊提示与安抚",
                "本地提醒并转人工",
            )
        ],
        widths=[Cm(3.34)] * 5,
        font_size=8.8,
        header_fill=DEEP_TEAL,
    )

    add_heading(doc, "3 3 典型服务闭环", 2)
    add_numbered(
        doc,
        [
            "用户说出目标，例如我先去卫生间，然后去挂号处，再去心内科。",
            "系统先检查停止和紧急求助关键词，再识别目的地与顺序。",
            "路线服务基于医院节点和道路连接组合完整路线，不用回答文本猜路线。",
            "界面同时展示每一站的位置说明、路径节点和地图高亮。",
            "用户点击开始导引。支持的实体路线下发给小车，其他路线由屏幕模拟。",
            "小车遇障碍停车并回传状态；到达后界面和语音同步提示。",
        ],
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    add_picture_with_alt(
        p.add_run(),
        ASSET_DIR / "yinban-interface-demo.png",
        width=Cm(16.2),
        alt="银伴适老界面：多目的地问答、路线图和模拟导引控制",
    )
    add_caption(doc, "图1  银伴适老界面和多目的地路线演示")

    add_heading(doc, "3 4 功能范围", 2)
    add_table(
        doc,
        ["级别", "内容", "进入条件"],
        [
            ("P0", "适老界面、语音文字按钮、规则问答、多站路线、模拟导引、单条实体路线、安全停车", "比赛前必须稳定"),
            ("P1", "第二条实体分支、可视化事件日志、本地灯光蜂鸣求助", "P0连续彩排通过后"),
            ("未来", "SLAM、多楼层、电梯联动、业务系统接口、个性化任务单", "完成真实场地试点与合规评估后"),
        ],
        widths=[Cm(2.0), Cm(10.0), Cm(4.7)],
    )

    add_section_break(doc)
    add_heading(doc, "4 原型成果与技术架构", 1)
    add_heading(doc, "4 1 已完成的软件成果", 2)
    add_bullets(
        doc,
        [
            "前端采用原生HTML CSS JavaScript，提供高对比大字号界面、语音输入、语音播报和文字兜底。",
            "后端采用FastAPI和Pydantic，提供对话、导航、机器人控制与状态接口。",
            "本地JSON知识库包含心内科、药房、卫生间、检验科和挂号处的位置与别名。",
            "意图路由支持导航、流程、求助、停止和未知五类结果，并优先处理安全关键词。",
            "路线规划使用节点和边计算道路组合，支持按自然语言顺序返回多个目的地。",
            "模拟机器人实现等待、移动、阻挡、恢复和到达状态，无硬件时可完整演示。",
            "硬件通信支持USB串口或经典蓝牙SPP虚拟串口，协议保持为易调试的单行文本。",
        ],
    )

    add_heading(doc, "4 2 总体架构", 2)
    add_table(
        doc,
        ["层级", "组件", "职责", "降级方式"],
        [
            ("交互层", "浏览器界面 Web Speech API", "接收语音文字按钮并展示路线状态", "语音失败时使用文字和按钮"),
            ("服务层", "FastAPI 对话与导航API", "组织问答路线和机器人命令", "保持本地运行不依赖外网"),
            ("知识层", "医院地点与流程JSON", "提供可审核的事实和别名", "未知问题转澄清或人工"),
            ("规划层", "图节点与路径组合", "生成多站点路线和实体路线编号", "不支持的实体路线转模拟"),
            ("通信层", "pyserial 蓝牙SPP USB", "发送命令并读取状态", "蓝牙失败改用USB或模拟"),
            ("执行层", "ESP32 状态机", "巡线避障停车和到达反馈", "本地STOP与超时独立生效"),
        ],
        widths=[Cm(2.1), Cm(4.1), Cm(6.3), Cm(4.2)],
        font_size=8.8,
    )

    add_heading(doc, "4 3 数据流", 2)
    add_numbered(
        doc,
        [
            "浏览器把用户文本发送到POST /api/dialogue。",
            "后端先执行安全规则，再解析意图与多个目的地。",
            "路线规划器依据routes.json连接节点，返回路线标签、坐标和实体路线编号。",
            "前端显示回答和路线，并在用户确认后调用POST /api/robot/start。",
            "机器人服务通过串口下发START STOP RESUME RESET，前端每500毫秒读取状态。",
            "到达、受阻、丢线或错误状态同步显示并触发相应播报。",
        ],
    )

    # Keep the protocol table from leaving only its repeated header at the
    # bottom of the preceding page in Word/PDF output.
    doc.add_page_break()
    add_heading(doc, "4 4 接口与协议", 2)
    add_table(
        doc,
        ["方向", "接口或消息", "作用"],
        [
            ("网页到后端", "POST /api/dialogue", "识别意图并返回回答与路线"),
            ("网页到后端", "POST /api/robot/start", "提交可执行的实体路线编号"),
            ("网页到后端", "POST /api/robot/stop", "最高优先级停止"),
            ("网页到后端", "GET /api/robot/status", "查询模式连接运动和错误状态"),
            ("电脑到小车", "START CARDIOLOGY", "启动心内科演示路线"),
            ("电脑到小车", "STOP RESUME RESET", "停止继续与复位"),
            ("小车到电脑", "MOVING BLOCKED ARRIVED", "回传移动受阻和到达状态"),
            ("小车到电脑", "LINE LOST ERROR", "回传丢线和传感器错误"),
        ],
        widths=[Cm(3.0), Cm(5.3), Cm(8.4)],
    )

    add_heading(doc, "4 5 软件质量", 2)
    add_para(
        doc,
        "当前自动测试覆盖五个测试文件，检查API正常与异常输入、目的地顺序、最短路径、协议解析和模拟机器人状态。2026年9月9日从项目目录执行测试，26项全部通过，同时出现2条第三方依赖弃用警告。警告不影响当前功能，但应在比赛后升级FastAPI和Starlette依赖并重新回归。",
    )

    add_section_break(doc)
    add_heading(doc, "5 硬件系统与安全设计", 1)
    add_heading(doc, "5 1 参赛硬件方案", 2)
    add_para(
        doc,
        "硬件已完成采购，计划书按500元上限管理。到货后不能直接照抄固件引脚，需要先核对卖家原理图和实物丝印。ESP32型号、传感器输出电平、电机驱动待机脚和电池电压是首要检查项。",
    )
    add_table(
        doc,
        ["模块", "推荐规格", "预算区间", "到货检查"],
        [
            ("主套件", "ESP32 WROOM 两轮车 五路巡线 超声波", "180至260元", "主控型号 焊接 资料 电池方案"),
            ("电机驱动", "TB6612FNG优先", "套件内", "STBY 电机供电 共地"),
            ("巡线", "五路可调 3.3V兼容", "15至30元", "输出电平 阈值 顺序"),
            ("测距", "HC SR04P或明确3.3V兼容型号", "8至20元", "Trig Echo电平和盲区"),
            ("供电", "带保护电池和匹配充电器", "50至90元", "极性 电压 充电与开关"),
            ("备件", "ESP32 扩展板 按键 杜邦线", "43至80元", "接口牢固和备用数量"),
            ("展台", "黑色胶带 白色KT板 固定材料", "48至95元", "路线对比度和固定强度"),
        ],
        widths=[Cm(2.4), Cm(5.2), Cm(2.5), Cm(6.6)],
        font_size=8.7,
    )
    add_para(
        doc,
        "预算区间来自赛前采购规格。套件已经包含的部件不得重复计算，实际结算以订单为准。若工具齐全，原型目标总价为300至400元；全部新购时应控制在480元以内并保留应急空间。",
    )

    add_heading(doc, "5 2 机器人状态机", 2)
    add_table(
        doc,
        ["状态", "进入条件", "执行动作", "退出条件"],
        [
            ("IDLE", "上电或复位", "电机停止 等待命令", "收到合法START"),
            ("FOLLOWING", "启动路线", "低速巡线并检查障碍", "到达 受阻 丢线 停止 超时"),
            ("BLOCKED", "连续测距低于阈值", "立即停车并回传距离", "障碍连续消失或收到STOP"),
            ("LINE LOST", "持续丢线超过阈值", "停车并等待人工处理", "复位后重新开始"),
            ("ARRIVED", "识别终点标记", "停车并回传到达", "收到RESET"),
            ("ERROR", "传感器或内部异常", "停车并回传错误", "排障后复位"),
        ],
        widths=[Cm(2.4), Cm(4.7), Cm(5.3), Cm(4.3)],
        font_size=8.8,
    )

    add_heading(doc, "5 3 安全优先级", 2)
    add_numbered(
        doc,
        [
            "STOP在任何状态下立即关闭两侧电机，不等待路线逻辑完成。",
            "启动后超过最长运行时间仍未到达时停车，防止失控持续运动。",
            "巡线信号短时波动允许纠偏，持续丢线必须停车。",
            "超声波读数连续多次低于阈值才判定障碍，障碍连续消失后才允许恢复。",
            "通信中断时小车执行本地停车，不能依赖网页端补发命令。",
            "比赛场地使用低速、宽弯和软质障碍物，电池充电远离观众区。",
        ],
    )

    add_heading(doc, "5 4 医疗与隐私边界", 2)
    add_bullets(
        doc,
        [
            "涉及胸痛、呼吸困难、晕倒等表达时，系统停止普通导航并提示立即联系现场医护或急救服务。",
            "流程回答使用预设文本，并提醒以医院当天通知为准。",
            "演示数据不包含真实姓名、身份证号、联系方式、病历、处方或人脸。",
            "日志只保留匿名测试文本、路线编号、状态和时间，比赛后可一键清除。",
            "未来接入医院系统前，需要完成数据最小化、授权、访问控制、日志审计和安全评估。",
        ],
    )

    add_section_break(doc)
    add_heading(doc, "6 创新与竞争分析", 1)
    add_heading(doc, "6 1 核心创新", 2)
    add_table(
        doc,
        ["创新点", "实现方式", "可验证证据"],
        [
            ("适老交互", "语音 大按钮 文字并存 高对比和少步骤", "现场由不同输入方式完成同一任务"),
            ("流程与路线融合", "回答下一步的同时生成可执行路线", "多站点问询与地图高亮一致"),
            ("低成本闭环", "普通电脑加ESP32巡线小车", "500元内硬件和本地软件联动"),
            ("确定性优先", "高频问题先走规则和本地知识库", "离线状态仍可完成主流程"),
            ("可降级设计", "蓝牙 USB 模拟机器人三层执行路径", "故障后30秒内切换演示"),
            ("边界可解释", "固定路线与未来自主导航明确分层", "答辩材料和界面均显示原型说明"),
        ],
        widths=[Cm(3.1), Cm(7.1), Cm(6.5)],
        font_size=8.8,
    )

    add_heading(doc, "6 2 替代方案比较", 2)
    add_table(
        doc,
        ["方案", "自然语言", "流程辅助", "实体带路", "成本与部署", "主要限制"],
        [
            ("人工导诊", "强", "强", "有限", "持续人员成本", "高峰期服务能力受限"),
            ("自助机", "弱", "中", "无", "已有基础设施", "需要用户学习操作"),
            ("手机地图", "中", "弱", "无", "边际成本低", "依赖手机和室内定位"),
            ("通用SLAM机器人", "强", "中", "强", "设备和部署成本高", "需地图运维与场地改造"),
            ("银伴原型", "强", "强", "固定路线", "硬件500元内", "当前只适合受控场地验证"),
        ],
        widths=[Cm(2.5), Cm(2.0), Cm(2.1), Cm(2.2), Cm(3.5), Cm(4.4)],
        font_size=8.3,
    )

    add_heading(doc, "6 3 可持续优势的形成条件", 2)
    add_para(
        doc,
        "原型本身尚未形成难以复制的技术壁垒。长期优势需要来自三类积累：经过真实用户测试的适老交互规范、可快速配置的医院流程与地图工具，以及在不同场地沉淀的安全运行数据。团队不应把通用语音识别或ESP32巡线算法包装为独有技术，而应围绕场景配置效率和可靠性建立知识产权。",
    )
    add_bullets(
        doc,
        [
            "申请软件著作权，覆盖场景问答、路线配置和机器人通信平台。",
            "对路线配置工具、异常恢复流程和适老交互组件保留版本记录与设计证据。",
            "引用开源项目时保留许可证和来源，不复制他人品牌、界面或机器人外观。",
            "注册银伴文字与图形商标前先进行近似检索。",
        ],
    )

    add_section_break(doc)
    add_heading(doc, "7 市场与商业路径", 1)
    add_heading(doc, "7 1 首批客户与使用地点", 2)
    add_para(
        doc,
        "项目采用机构客户优先的路径。医院、社区卫生服务中心、体检中心和养老服务机构可以购买或租用设备，并配置本机构地图与流程。老年用户不直接承担硬件购买和维护。首个试点应选择单栋门诊楼、单层或两点之间的固定路线，先验证服务价值，再扩展地图和接口。",
    )
    add_table(
        doc,
        ["客户", "需求", "切入产品", "决策关注"],
        [
            ("医院门诊部", "降低重复问询 改善体验", "大厅问询和固定路线导引", "安全 可靠 运维 合规"),
            ("社区卫生中心", "人员少 服务对象老龄化", "桌面终端和简化移动底盘", "成本 易用 本地维护"),
            ("体检中心", "多站点顺序复杂", "个性化检查路线清单", "准时率 路线灵活性"),
            ("养老服务机构", "陪诊和外出服务", "陪诊任务提示与人员协作", "数据授权 责任边界"),
            ("交通枢纽", "问询与固定区域导引", "机场车站场景包", "大客流 多语言 无障碍"),
        ],
        widths=[Cm(3.1), Cm(4.7), Cm(5.0), Cm(3.9)],
        font_size=8.8,
    )

    add_heading(doc, "7 2 商业模式", 2)
    add_para(
        doc,
        "银伴采用机构付费、老年用户免费使用的模式。首批付费方定位为社区卫生服务中心、医院门诊部和体检中心。比赛原型不直接出售；对外先交付30天受控试用，再转为标准设备采购、场景部署和年度软件服务。以下价格作为首轮执行报价，而不是待定设想。",
    )
    add_table(
        doc,
        ["收费项目", "执行价格", "包含内容", "验收与收费规则"],
        [
            ("30天验证包", "2980元 每点位", "1台设备 1个楼层 不超过10个地点 1条实体路线 培训与验证报告", "完成部署后收费；90天内采购设备可抵扣1500元"),
            ("标准设备", "12800元 每台", "受控区域移动底盘 交互终端 充电装置 1年基础保修", "合同签订付50% 到货部署付40% 30天验收后付10%"),
            ("场景部署", "3000元 每点位", "首层地图 10个地点 1条实体路线 知识库初始化与一次培训", "新增楼层1500元 新增实体路线800元 超出10个地点后每个100元"),
            ("年度软件服务", "1200元 每台每年", "版本更新 知识库维护 日志导出与远程支持", "首年包含在设备价内 第二年起按年续费"),
        ],
        widths=[Cm(3.0), Cm(2.6), Cm(6.5), Cm(4.6)],
        font_size=8.2,
    )
    add_para(
        doc,
        "首个标准点位配置1台设备时，第一年合同金额为15800元，包括12800元设备和3000元场景部署。内部成本控制目标为：产品化硬件与装配不超过5500元，地图部署与差旅不超过1500元，首年保修准备金不超过1000元，单点位直接成本合计不超过8000元，目标贡献毛额不低于7800元，约占合同金额49%。该比例未扣除销售、管理和研发费用，只用于控制单项目交付成本。",
    )
    add_para(
        doc,
        "标准合同的验收条件统一为：预设地点问答15条全部正确；指定路线20次运行至少18次成功；急停、障碍、丢线、超时和断连测试全部触发安全停车；部署清单、培训记录和故障处理说明完成交接。医院信息系统、电梯控制和真实患者数据接口不包含在首轮标准产品中，未经合规和安全评估不得承诺。",
    )

    add_heading(doc, "7 3 试点路径", 2)
    add_numbered(
        doc,
        [
            "完成首轮5至8名老年用户或家属访谈，并补充1至2名服务人员访谈，确认高频任务、拒绝原因和禁用表述。",
            "在校园或模拟门诊搭建单条路线，完成30次封闭环境运行。",
            "与一家社区卫生中心或医院志愿服务部门讨论非临床、非联网的桌面试用。",
            "试点只记录匿名任务完成数据，人工人员全程可停止设备。",
            "达到安全和可用性门槛后，再增加第二条路线和远程运维能力。",
        ],
    )

    add_heading(doc, "7 3 1 分阶段转化门槛", 3)
    add_table(
        doc,
        ["阶段", "交付重点", "进入下一阶段的条件"],
        [
            ("概念验证", "适老交互 路线展示 模拟机器人闭环", "软件测试通过，核心问法和多站顺序正确"),
            ("单楼层准备", "一条实体路线 安全状态机 地图版本", "20次彩排至少18次成功，急停和降级均有效"),
            ("受控场地试点", "匿名用户测试 人工接管 运维记录", "获得场地方书面同意，任务完成率和接管率可记录"),
            ("多路线验证", "第二条路线 地图配置 异常恢复", "路线更新责任明确，连续运行达到既定门槛"),
            ("产品化评估", "结构加固 多设备管理 合规与成本", "存在继续试用或采购意向，部署和维护成本可核算"),
        ],
        widths=[Cm(3.0), Cm(6.2), Cm(7.5)],
        font_size=8.6,
    )
    add_para(
        doc,
        "每个试点地图必须标明采集日期、适用楼层、节点来源、审核人和更新责任人。场地道路或科室位置变化后，应停用旧版本，复核路线后再恢复导引。",
    )

    add_heading(doc, "7 4 首轮客户获取与商业验证", 2)
    add_para(
        doc,
        "比赛阶段不以虚构销售额证明市场，而是用可核验的客户接触和试用意向验证商业路径。团队执行以下最低目标，并将未达到的项目如实记录为未完成。",
    )
    add_table(
        doc,
        ["任务", "数量目标", "完成证据", "主责岗位"],
        [
            ("建立机构名单", "8家", "单位名称 联系部门 联系日期和下一步", "调研与商务"),
            ("完成决策人员访谈", "至少3人", "匿名纪要 采购关注点 主要反对意见", "项目负责人"),
            ("发送验证包方案", "至少3份", "统一报价 服务边界和30天试用清单", "调研与商务"),
            ("取得书面试用意向", "至少1份", "邮件 回函或非约束性意向说明", "项目负责人"),
            ("形成价格复盘", "1份", "客户反馈 实际成本和价格调整理由", "项目负责人"),
        ],
        widths=[Cm(4.0), Cm(2.4), Cm(7.1), Cm(3.2)],
        font_size=8.5,
    )

    add_section_break(doc)
    add_heading(doc, "8 实施与资源计划", 1)
    add_heading(doc, "8 1 硬件到货后的七天计划", 2)
    add_table(
        doc,
        ["日期", "重点任务", "当日完成条件"],
        [
            ("Day 1", "清点器件 核对电压与引脚 固定模块", "形成实物接线表 上电无异常"),
            ("Day 2", "单测左右电机 巡线传感器 超声波 按键", "每个模块能独立读取或控制"),
            ("Day 3", "标定黑白阈值和低速巡线参数", "直线和缓弯连续成功5次"),
            ("Day 4", "实现障碍停车 丢线停车 超时和本地STOP", "五类安全测试全部触发停车"),
            ("Day 5", "接入USB串口 再测试蓝牙SPP", "网页命令与状态回传形成闭环"),
            ("Day 6", "搭建医院路线 联调界面 地图和小车", "完整流程连续成功10次"),
            ("Day 7", "冻结功能 完成20次彩排和备用视频", "成功至少18次 演示少于2分钟"),
        ],
        widths=[Cm(2.1), Cm(7.4), Cm(7.2)],
        font_size=8.8,
    )

    add_heading(doc, "8 2 赛前里程碑与交付物", 2)
    add_para(
        doc,
        "本节以正式答辩或现场比赛日期为T，按交付物和验收条件倒排准备工作。若准备时间不足21天，可合并相邻阶段，但不能删除对应验收条件。",
    )
    add_table(
        doc,
        ["时间", "必须完成的工作", "交付物", "通过条件"],
        [
            ("T减21至14天", "冻结参赛范围 清点硬件 建立访谈和机构名单", "实物清单 接线方案 任务看板", "不再增加非核心功能"),
            ("T减14至7天", "完成单模块测试 安全停车 软件与小车闭环", "测试日志 接线照片 第一版完整视频", "完整流程连续成功10次"),
            ("T减7至3天", "完成用户访谈 商业材料 计划书和答辩稿", "匿名访谈记录 报价单 计划书 PPT", "证据编号与正文结论一致"),
            ("T减3至1天", "冻结版本 完成20次彩排和故障切换", "最终软件包 彩排表 备用视频", "至少18次成功 故障30秒内降级"),
            ("T减1天", "双设备拷贝 离线检查 展台装箱 人员走位", "参赛材料清单 两份启动包", "断网状态下仍可完成核心演示"),
            ("比赛当天", "提前启动 复测主路线 按两分钟脚本演示", "现场记录 评委问题清单", "演示后立即保存日志并复盘"),
        ],
        widths=[Cm(3.0), Cm(6.1), Cm(4.5), Cm(3.1)],
        font_size=8.2,
    )

    add_heading(doc, "8 3 团队岗位配置", 2)
    add_para(
        doc,
        "项目按五个主责岗位推进。成员可以兼任岗位，但每项交付只能有一名最终负责人，并必须指定另一名成员复核。团队成员姓名在报名表中与下列岗位一一对应，不能使用全员负责替代明确分工。",
    )
    add_table(
        doc,
        ["主责岗位", "必须完成的工作", "截止时间", "验收与交接"],
        [
            ("项目负责人", "锁定范围 排期 决策 机构沟通 主持答辩", "T减14天定范围 T减3天批准冻结", "任务看板无遗漏 关键数据有证据 可回答能力边界"),
            ("软件负责人", "前后端 意图识别 路线规划 机器人接口 自动测试", "T减7天完成功能 T减3天冻结", "测试全部通过 离线启动成功 多站顺序与地图一致"),
            ("硬件负责人", "接线 固件 巡线 避障 电源和本地急停", "到货后7天完成 T减3天冻结", "五类安全测试通过 20次彩排有记录 接线可复现"),
            ("调研与商务", "用户访谈 机构名单 验证包报价 试用意向", "T减7天完成访谈 T减3天汇总", "5至8名用户及1至2名服务人员记录 商业目标逐项留证"),
            ("材料与演示", "计划书 PPT 两分钟脚本 视频 展台与备份包", "T减3天定稿 T减1天装箱", "主备文件均可打开 演示不超时 现场切换路径明确"),
        ],
        widths=[Cm(3.0), Cm(6.5), Cm(3.2), Cm(4.0)],
        font_size=8.0,
    )
    add_para(
        doc,
        "若团队只有两至三人，项目负责人可兼任调研与商务，软件负责人可兼任材料与演示；硬件负责人在带电和运动测试时不得同时担任唯一安全观察者，必须由另一名成员掌握急停。",
    )

    add_heading(doc, "8 3 1 团队协作与版本冻结机制", 3)
    add_bullets(
        doc,
        [
            "代码、固件和文档统一使用Git保存，演示版本设置明确标签；硬件联调Day 6或T减3天两者较早的时间冻结功能，只修复影响主流程和安全的问题。",
            "每项关键功能同时准备正常流程、可预期的失败流程和证明材料，避免只展示成功画面。",
            "测试数据只有在保留日志、截图或视频后才能写入计划书和答辩稿，预测值不得写成实测结果。",
            "软件、硬件和演示操作至少由另一名成员交叉复核，确保关键流程不依赖单个人员。",
            "每次彩排只记录一个主要失败原因，并指定责任人、修复期限和复测结果。",
        ],
    )

    add_heading(doc, "8 4 资源预算", 2)
    add_para(
        doc,
        "参赛阶段总预算上限固定为500元。硬件已经购买后，团队按订单和实物重新登记实际金额；计划书先以以下分类上限控制支出，节余统一转入应急额度，任何类别超支都必须从非安全类支出中等额削减。",
    )
    add_table(
        doc,
        ["预算类别", "金额上限", "必须覆盖", "支出控制"],
        [
            ("小车 主控与驱动", "260元", "底盘 电机 ESP32 驱动板 基础线材", "套件已有部件不得重复计价"),
            ("供电与安全", "80元", "保护电池 匹配充电器 电源开关 急停", "安全件不得为外观材料让预算"),
            ("传感器与备件", "70元", "巡线 测距 按键 杜邦线和必要备用件", "同类备用件原则上不超过1套"),
            ("路线与展台材料", "60元", "胶带 KT板 固定件 标识打印", "优先复用现有工具和材料"),
            ("应急额度", "30元", "临时替换线材 接头或固定件", "非故障不动用 使用后记录原因"),
        ],
        widths=[Cm(3.4), Cm(2.5), Cm(6.7), Cm(4.1)],
        font_size=8.5,
    )

    add_heading(doc, "8 4 1 经费使用与凭证管理原则", 3)
    add_bullets(
        doc,
        [
            "500元参赛预算优先保障电源安全、急停、固定件和必要备用件，其次才是外观装饰。",
            "保存订单、发票或付款记录，并在采购表中记录单价、数量、用途、到货状态和实物型号。",
            "套件中已包含的主控、电机、传感器或线材应标记，避免重复采购和重复计入预算。",
            "参赛样机预算与医院试点预算分开核算，后续规划金额不得解释为当前已经投入或已经获得的资金。",
            "具体型号和电气参数以到货实物、芯片丝印和说明书复核结果为准，不使用商家简称替代技术规格。",
        ],
    )

    add_section_break(doc)
    add_heading(doc, "9 验证方案与风险管理", 1)
    add_heading(doc, "9 1 验证指标", 2)
    add_table(
        doc,
        ["指标", "当前状态", "比赛目标", "记录方式"],
        [
            ("自动测试", "26项通过", "保持全部通过", "pytest报告"),
            ("预设问法", "软件已覆盖", "15条全部正确", "输入与输出记录"),
            ("多站点顺序", "已实现", "指定测试句顺序一致", "路线节点截图"),
            ("障碍停车", "固件待联调", "5次全部停车", "距离与状态日志"),
            ("丢线与超时", "协议已设计", "触发后立即停车", "串口日志和视频"),
            ("完整演示", "软件闭环完成", "20次至少18次成功", "彩排表"),
            ("故障降级", "模拟模式完成", "30秒内恢复演示", "计时记录"),
            ("演示时长", "待彩排", "主流程不超过2分钟", "视频时间轴"),
        ],
        widths=[Cm(3.1), Cm(3.5), Cm(5.0), Cm(5.1)],
        font_size=8.7,
    )

    add_heading(doc, "9 2 硬件测试顺序", 2)
    add_numbered(
        doc,
        [
            "断开电机负载，确认主控供电、电压和串口输出。",
            "架空车轮测试左右电机方向与STOP，错误时只修改引脚映射。",
            "手持巡线模块越过黑白边界，记录五路原始状态和阈值。",
            "低速测试直线，再增加缓弯，不同时调速度、阈值和转向三个变量。",
            "使用软纸盒测试障碍出现、持续和移除，记录误报与漏报。",
            "拔出通信线或关闭蓝牙验证本地停车，不让软件界面代替硬件安全。",
            "最后进行全流程彩排，每次只记录一个主要失败原因。",
        ],
    )

    add_heading(doc, "9 3 现场演示脚本", 2)
    add_table(
        doc,
        ["时间", "操作", "讲解重点"],
        [
            ("0至20秒", "说明老人就医中的信息和流程问题", "项目服务于数字化困难而非替代医护"),
            ("20至45秒", "说出我要去心内科或多站点请求", "本地知识库和路线规划不依赖外网"),
            ("45至80秒", "点击开始导引 小车沿线运行", "网页 地图 通信和实体执行形成闭环"),
            ("80至105秒", "放置软障碍物后移除", "障碍停车和恢复状态同步"),
            ("105至120秒", "到达并说明产品边界", "当前是固定环境概念验证"),
        ],
        widths=[Cm(2.8), Cm(6.1), Cm(7.8)],
    )

    add_heading(doc, "9 4 降级预案", 2)
    add_table(
        doc,
        ["故障", "现场动作", "继续展示的价值"],
        [
            ("语音识别失败", "立即使用大按钮或文字输入", "适老多入口和问答路线仍可展示"),
            ("蓝牙失败", "切换USB串口", "保留真实小车控制"),
            ("小车故障", "切换模拟模式并播放备用视频", "保留软件闭环和安全逻辑说明"),
            ("网络不可用", "继续使用本地知识库和页面", "证明核心流程不依赖云服务"),
            ("页面异常", "使用启动器重启本地服务", "不在现场安装或升级依赖"),
        ],
        widths=[Cm(3.2), Cm(6.0), Cm(7.5)],
    )

    add_heading(doc, "9 5 风险矩阵", 2)
    add_table(
        doc,
        ["风险", "可能性", "影响", "预防与触发条件", "备用方案"],
        [
            ("实物型号与预期不同", "中", "高", "到货先核对主控和电平", "按实物重写引脚 不强刷固件"),
            ("巡线在弯道丢线", "中", "高", "低速 宽弯 单变量标定", "缩短路线或降低速度"),
            ("电机干扰重启", "中", "高", "检查供电 共地和线材固定", "更换电池或分离供电"),
            ("超声波误报", "中", "中", "连续多次读数再判定", "现场人工STOP"),
            ("语音受噪声影响", "高", "中", "测试麦克风并显示识别文本", "按钮与文字输入"),
            ("蓝牙配对失败", "中", "中", "提前绑定COM口并保留截图", "USB串口"),
            ("能力被误解为医疗诊断", "低", "高", "界面和答辩明确非诊疗边界", "删除可能引起误解的表述"),
            ("日志包含个人信息", "低", "高", "禁止真实患者测试 数据最小化", "关闭日志并清除测试数据"),
            ("功能持续增加", "中", "中", "Day 6后冻结功能", "只保留一条实体路线"),
        ],
        widths=[Cm(3.0), Cm(1.7), Cm(1.7), Cm(6.0), Cm(4.3)],
        font_size=8.1,
    )

    add_heading(doc, "9 6 证据与版本管理", 2)
    add_para(
        doc,
        "项目结论按照可追溯原则管理。计划书中的已完成状态、测试数据和现场表现必须能回到原始记录；无法追溯的数据只作为假设，不作为成果。",
    )
    add_table(
        doc,
        ["证据类别", "最少保留内容", "使用规则"],
        [
            ("软件", "Git提交 测试输出 关键页面截图 版本标签", "截图对应明确版本，测试失败不得删除"),
            ("硬件", "接线照片 实物型号 传感器参数 串口日志", "记录测试条件，型号以实物核验为准"),
            ("演示", "完整未剪辑视频 彩排表 故障恢复记录", "备用视频须标明录制时间，不冒充现场运行"),
            ("用户调研", "匿名访谈记录 样本说明 反对意见 失败案例", "取得知情同意，不保存姓名和诊疗信息"),
            ("市场与政策", "原始链接 发布机构 日期 计算过程", "区分事实、推算和团队判断"),
        ],
        widths=[Cm(3.2), Cm(7.0), Cm(6.5)],
        font_size=8.6,
    )
    add_para(
        doc,
        "建议采用日期加版本号命名证据文件，并建立证据索引表，至少包含证据编号、关联结论、负责人、生成日期和存放位置。提交材料前复核隐私、开源许可证和第三方素材来源。",
    )

    add_section_break(doc)
    add_heading(doc, "10 社会价值与发展边界", 1)
    add_heading(doc, "10 1 社会价值", 2)
    add_para(
        doc,
        "银伴把公共空间中的数字化服务转换为更直接的语音、文字和实体行动。它的价值不在于要求老年人学习另一套复杂系统，而在于让同一任务可以通过说话、按钮或文字完成，并在机器人无法执行时明确转向人工。",
    )
    add_bullets(
        doc,
        [
            "降低第一次到院和多科室就诊时的信息负担。",
            "为视力下降或不熟悉自助机的用户提供更大的操作目标和语音反馈。",
            "分担导诊人员的高频重复问询，但保留人工服务和紧急处置。",
            "以低成本原型验证公共服务机器人是否真正改善任务完成率。",
            "为机场、车站和政务大厅等场景积累可迁移的适老交互方法。",
        ],
    )

    add_heading(doc, "10 2 扩展路线", 2)
    add_table(
        doc,
        ["阶段", "场景", "新增能力", "前置条件"],
        [
            ("当前", "模拟医院", "问询 多站路线 单条实体巡线", "比赛环境安全验证"),
            ("下一步", "社区卫生中心", "地图配置 远程运维 桌面或移动终端", "机构合作和匿名用户测试"),
            ("中期", "大型医院封闭区域", "多设备管理 电梯口交接 任务工单", "合规评估和现场责任机制"),
            ("远期", "机场 车站 政务大厅", "多语言 场景知识包 动态导航", "新场景验证和硬件升级"),
        ],
        widths=[Cm(2.4), Cm(3.5), Cm(6.4), Cm(4.4)],
        font_size=8.8,
    )

    add_heading(doc, "10 3 阶段结论", 2)
    add_para(
        doc,
        "截至2026年9月10日，银伴软件演示版本已经完成，26项自动化测试全部通过。五类地点问询、多目的地顺序解析、道路节点组合、路线动画、模拟机器人状态以及启动停止协议均有可复核结果。因此，软件端已经达到本计划书设定的参赛演示验收线，可以独立完成两分钟核心流程演示。",
    )
    add_para(
        doc,
        "实体机器人尚未完成到货后的型号核验、接线、巡线标定、障碍停车和整机连续运行测试。因此，项目当前没有通过实体机器人整机验收，不能宣称已经具备真实医院自主导引能力。现阶段产品性质确定为医院固定环境参赛原型，不属于可直接交付医院使用的正式产品。",
    )
    add_para(
        doc,
        "比赛是否采用实体小车作为主展示方式，按统一门槛决定：指定路线运行20次至少成功18次；5次障碍测试全部停车；人工STOP、丢线、通信中断和运行超时均能触发本地停机；完整演示不超过120秒；故障后30秒内能够切换到模拟模式。五项条件全部满足时，比赛采用软件与实物闭环演示；任一条件未满足时，实体小车退出主流程，正式演示改用软件模拟，小车仅作静态展示，备用视频必须明确标注为预录证据。该决定在比赛前一天完成并冻结，现场不再增加功能或修改安全参数。",
    )

    add_heading(doc, "附录 A 最小功能清单", 1)
    add_table(
        doc,
        ["功能", "状态", "验收方法"],
        [
            ("大按钮 文字 语音输入", "软件已完成", "三种方式均能提交同一目的地"),
            ("五类地点问询", "软件已完成", "心内科 药房 卫生间 检验科 挂号处均返回位置"),
            ("多目的地顺序", "软件已完成", "路线顺序与用户表达一致"),
            ("真实道路组合", "软件已完成", "路线只通过已定义边连接"),
            ("屏幕模拟导引", "软件已完成", "所有已规划路线均可启动动画"),
            ("模拟机器人状态", "软件已完成", "移动 受阻 恢复 到达可展示"),
            ("单条实体巡线", "待硬件联调", "门诊大厅到心内科连续5次"),
            ("障碍停车", "待硬件联调", "5次障碍测试全部停车"),
            ("STOP RESUME RESET", "协议已完成 待实物验证", "运行中命令均得到正确状态"),
            ("丢线 通信中断 超时停车", "固件已设计 待实物验证", "逐项制造故障并记录"),
            ("到达播报", "软件已完成 待联调", "小车回传ARRIVED后自动播报"),
            ("人工求助", "本地提醒已完成", "不得宣称已通知真实医护"),
        ],
        widths=[Cm(5.0), Cm(4.3), Cm(7.4)],
        font_size=8.7,
    )

    add_heading(doc, "附录 B 软件目录", 1)
    directory_text = (
        "yinban-mvp\n"
        "  app\n"
        "    api          对话 导航与机器人接口\n"
        "    core         意图 路线 对话与安全规则\n"
        "    data         医院知识和地图数据\n"
        "    services     模拟机器人 串口与协议\n"
        "    web          适老界面 地图和前端逻辑\n"
        "  firmware\n"
        "    yinban_car   ESP32 Arduino固件\n"
        "  tests          自动测试\n"
        "  docs           接线 采购 演示与故障文档\n"
        "  run.ps1        PowerShell启动脚本\n"
        "  start-yinban.cmd  Windows双击启动器"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(directory_text)
    set_run_font(run, size=9.2, name="Cascadia Mono")

    add_heading(doc, "附录 C 品牌文件", 1)
    brand_table = doc.add_table(rows=1, cols=2)
    brand_table.autofit = False
    remove_table_borders(brand_table)
    left, right = brand_table.rows[0].cells
    left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(
        lp.add_run(),
        ASSET_DIR / "yinban-logo-square.png",
        width=Cm(6.2),
        alt="银伴报名方形Logo",
    )
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(
        rp.add_run(),
        ASSET_DIR / "yinban-logo-mark-transparent.png",
        width=Cm(6.2),
        alt="银伴透明背景机器人图标",
    )
    add_caption(doc, "图2  报名方形版与透明机器人图标版")
    add_para(
        doc,
        "Logo使用墨绿、暖金、奶白和银灰四类颜色。机器人顶部的路线环和定位标记表示导引，胸前的环抱路径表示陪伴。形象参考了圆润服务机器人的亲和感，但没有复制参考产品的桶帽、脸型或具体结构。",
    )

    add_heading(doc, "附录 D 资料来源", 1)
    sources = [
        (
            "[1] 国家统计局 中华人民共和国2025年国民经济和社会发展统计公报 2026年2月28日",
            "https://www.stats.gov.cn/xxgk/sjfb/tjgb2020/202602/t20260228_1962662.html",
        ),
        (
            "[2] 国务院办公厅 关于切实解决老年人运用智能技术困难的实施方案 国办发2020 45号",
            "https://www.nhc.gov.cn/bgt/gwywj2/202011/b51828e5adac4dbc92fbaa8d26474802.shtml",
        ),
        (
            "[3] 国家卫生健康委 国家中医药局 改善就医感受提升患者体验主题活动方案 2023至2025年",
            "https://www.nhc.gov.cn/wjw/c100375/202305/bfa59db84f1041c4b50bb859a3b76f39.shtml",
        ),
        (
            "[4] 中华人民共和国无障碍环境建设法 2023年6月28日通过",
            "https://gxca.miit.gov.cn/zwgk/zcwj/flfg/art/2024/art_340a68cb6a8e45ae81c121630dd37346.html",
        ),
        (
            "[5] 国家卫生健康委 对十四届全国人大三次会议第7409号建议的答复 2025年8月",
            "https://www.nhc.gov.cn/wjw/jiany/202508/c66483762ee2451fa550847bf28e0e89.shtml",
        ),
        (
            "[6] 银伴项目代码仓库",
            "https://github.com/zhaiqianni/yinban-mvp",
        ),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        set_run_font(p.add_run(label + "  "), size=9.5)
        add_hyperlink(p, url, url)
    add_para(doc, "公开资料访问日期为2026年9月10日。代码仓库提交参赛材料前应按赛事要求设置公开性或评审访问权限。", italic=True)

    doc.add_page_break()
    add_heading(doc, "附录 E 答辩证据清单", 1)
    add_table(
        doc,
        ["检查项", "答辩前完成要求", "状态"],
        [
            ("软件版本", "冻结可运行版本并记录Git标签和提交号", "□"),
            ("自动测试", "保存完整测试输出，所有核心测试通过", "□"),
            ("多站路线", "保留指定问法、路线节点和地图高亮截图", "□"),
            ("硬件核验", "保存实物照片、型号、电压、引脚和接线表", "□"),
            ("安全测试", "急停、障碍、丢线、超时和断连逐项留档", "□"),
            ("完整彩排", "20次彩排至少18次成功并记录失败原因", "□"),
            ("备用方案", "USB、模拟模式和备用视频均能在30秒内切换", "□"),
            ("用户证据", "完成匿名访谈并保留原始记录和反对意见", "□"),
            ("来源审查", "政策、数字、图片和开源代码均有可查来源", "□"),
            ("边界表述", "不宣称诊断、真实医院部署或未验证能力", "□"),
        ],
        widths=[Cm(3.5), Cm(11.3), Cm(1.9)],
        font_size=8.7,
    )

    add_heading(doc, "附录 F 用户访谈提纲", 1)
    add_para(
        doc,
        "访谈对象以老年就医者和陪同家属为主。访谈开始前说明用途、匿名方式和自愿退出权，不记录姓名、身份证、病历、诊断结果或联系方式。",
    )
    add_table(
        doc,
        ["序号", "访谈问题", "记录重点"],
        [
            ("1", "请回忆最近一次到医院办事的完整过程。", "场景 时间顺序 是否有人陪同"),
            ("2", "其中最困难或最容易走错的是哪一步？", "地点 流程 信息理解 操作障碍"),
            ("3", "遇到困难时，您通常向谁求助或使用什么办法？", "人工窗口 家属 手机 标识 放弃任务"),
            ("4", "如果一次要去多个地点，您怎样记住顺序？", "纸条 截图 口头提醒 是否会临时改路线"),
            ("5", "使用语音、大按钮和文字输入时，您更倾向哪一种？为什么？", "选择原因 噪声 视力 识字 隐私"),
            ("6", "如果机器人低速带路，什么情况会让您愿意或拒绝跟随？", "速度 距离 外观 人员陪同 安全顾虑"),
            ("7", "机器人停止、迷路或无法回答时，您希望它怎样处理？", "人工接管 求助入口 提示方式 可接受等待"),
        ],
        widths=[Cm(1.6), Cm(8.8), Cm(6.3)],
        font_size=8.7,
    )
    add_para(
        doc,
        "访谈结束后按统一编号保存原始记录，并把每条需求标记为已验证、部分验证或未验证。少数意见和拒绝意见同样保留，不能只汇总支持项目的回答。",
    )

    add_heading(doc, "附录 G 两分钟演示操作清单", 1)
    add_table(
        doc,
        ["时间", "演示动作", "操作检查与异常分支"],
        [
            ("演示前", "启动本地服务 连接小车 清空旧任务", "确认电量 路线 障碍物 音量 STOP和USB备用"),
            ("0至20秒", "说明老年人多站就医中的信息负担", "不扩展到诊断或陌生环境自主导航"),
            ("20至45秒", "说出单站或多站请求并核对识别文本", "语音失败立即改用大按钮或文字"),
            ("45至80秒", "展示完整路线并点击开始导引", "图标必须沿道路节点移动；蓝牙失败切换USB"),
            ("80至105秒", "放置软障碍并展示停车和恢复", "状态未回传时立即STOP，不反复重启电机"),
            ("105至120秒", "到达播报并说明能力边界", "小车故障则切换模拟模式，备用视频只作证据说明"),
            ("结束后", "停止电机 保存日志 记录本次失败原因", "不在现场临时升级依赖或增加功能"),
        ],
        widths=[Cm(2.4), Cm(6.6), Cm(7.7)],
        font_size=8.5,
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
