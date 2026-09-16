from __future__ import annotations

from hashlib import sha256
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = Path(
    r"D:\Wchat document\xwechat_files\wxid_xgizti2h0n9a12_8955\msg\file\2026-09\商业计划书模板.docx"
)
OUTPUT_PATH = BASE_DIR / "银伴商业计划书（模板版）.docx"
ASSET_DIR = BASE_DIR / "assets"
PROJECT_ROOT = BASE_DIR.parents[1]
SCREENSHOT_DIR = PROJECT_ROOT / "output" / "playwright" / "plan-evidence"
EXPECTED_TEMPLATE_SHA256 = "133a563b4a4008938310e4224795f1c1cc7be57620ee12288f6978e3e753a5f5"

FONT_CN = "Microsoft YaHei"
FONT_COVER = "FZSmallBiaoSong-B06S"
FONT_EN = "Arial"
BLACK = "000000"
BLUE = "2E75B6"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F5F5"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


SECTION_ENRICHMENTS = {
    "第一章 执行概要": [
        "本章从使用需求、产品形态、市场验证、收入逻辑、团队执行和风险门槛六个方面概括项目。所有结论均区分已完成的软件成果、到货后才能验证的实车能力以及需要机构合作才能开展的试点任务，以保证申报材料与当前证据一致。",
    ],
    "1.1 项目概况": [
        "银伴选择医院作为首个应用场景，是因为院内任务同时具有地点密集、顺序连续、楼层切换频繁和错误成本较高等特征。项目把一次就医拆解为可确认的目的地序列和可追踪的路线节点，让用户知道现在在哪里、下一步去哪里、何时需要乘梯以及遇到困难如何停止或求助。",
    ],
    "1.2 产品介绍": [
        "产品设计强调低认知负担和可解释执行。用户不必学习复杂菜单，语音、文字和常用服务按钮都进入同一任务处理链；系统返回地点事实、完整路线和当前状态，而不是只给出一个科室位置。软硬件采用统一命令与状态协议，便于在模拟机器人、USB串口和蓝牙链路之间切换。",
    ],
    "1.3 市场分析": [
        "项目将需求规模与可支付市场分开判断。老龄人口和诊疗人次用于说明问题的长期性与高频性，实际商业判断则以单点位任务量、重复问询占比、人工接管次数、部署维护成本和采购意向为依据，避免把宏观统计直接换算成未经验证的销售预测。",
    ],
    "1.4 商业模式": [
        "验证包、设备、部署和年度服务分别对应决策试用、硬件交付、场景配置和持续维护四类价值。该拆分使机构可以先用较低成本验证场景，再决定是否采购；同时也便于团队单独核算硬件、地图配置、培训和售后的真实成本。",
    ],
    "1.5 组织管理": [
        "团队实行主责人交付、另一成员复核的工作机制。需求和商业口径由项目负责人冻结，软件版本由软件负责人冻结，带电与运动测试由硬件负责人组织；涉及路线安全、数据真实性和对外表述的材料必须经过交叉确认，减少单人判断造成的技术或表述偏差。",
    ],
    "1.6 风险管理": [
        "风险管理不以笼统的注意事项结束，而是为每项高影响风险设置触发条件、责任人和降级动作。任何影响制动、通信或路线一致性的问题优先停止实体演示；语音和蓝牙故障允许切换到文字、按钮或USB；若整车未通过验收，比赛主流程只使用软件模拟并明确标注。",
    ],
    "第二章 项目概况": [
        "本章说明银伴要解决的核心问题、已经形成的成果、团队执行结构和项目价值取向。项目的专业性不以功能数量衡量，而以任务是否闭环、路线是否可解释、异常是否可停止、结论是否有证据为主要判断标准。",
    ],
    "2.1 项目介绍": [
        "综合来看，银伴的优势集中在五个方面：面向老年用户的低学习成本交互、能够保持目的地顺序的分楼层路线、无需外网的本地确定性问答、网页与ESP32协同的双层安全机制，以及500元以内即可搭建的比赛原型。创新并非来自单一算法，而是来自对医院连续任务、跨楼层过渡和人机交接的系统化组合。",
    ],
    "2.2 当前成果": [
        "现阶段成果已经形成可运行、可测试和可演示的软件闭环。网页端能够接收多目的地需求、返回地点和电梯过渡信息、按楼层显示路线并驱动模拟机器人状态；代码端保留自动测试和版本记录。实物成果仍以到货后的型号核验与联调结果为准，两类成果在材料中分别表述。",
    ],
    "2.3 组织架构": [
        "三人团队采用轻量级责任矩阵管理：每项任务只能有一名主责人，但关键结果必须有明确复核人。软件、硬件和商业材料分别建立版本号、验收记录和问题清单；比赛前冻结主流程后，只处理安全缺陷、阻断性错误和证据不一致问题。",
    ],
    "2.4 项目宗旨": [
        "银伴不以减少所有人工岗位为目标，而是承担高频、重复且规则明确的信息与固定路线任务，把现场人员保留给复杂咨询、突发情况和需要情绪安抚的服务。该定位既符合适老服务需要人工兜底的现实，也降低原型阶段过度自动化带来的安全风险。",
    ],
    "第三章 市场分析": [
        "本章采用从宏观背景到具体任务、再到采购约束的分析路径。宏观数据说明适老服务需求具有持续性，用户旅程分析说明问题发生在哪里，竞品比较和验证包设计则回答银伴为什么值得被机构试用。",
    ],
    "3.1 行业背景概述": [
        "行业机会来自数字化服务效率与部分老年用户适应能力之间的落差，而不是简单来自人口数量增长。医院的线上线下一体化程度越高，越需要提供语音、大字、人工协助和连续路线说明等替代入口，使效率提升不会转化为新的使用障碍。",
    ],
    "3.2 目标客户": [
        "用户、受益者和采购者并不完全相同。老年就医者关注是否容易理解和安全，家属关注多任务安排是否清楚，导诊人员关注能否减少重复解释，医院管理者则关注责任边界、维护成本和数据合规。产品验证需要同时覆盖这四类评价，而不能只采集终端用户好评。",
    ],
    "3.3 目标客户面临的问题": [
        "目标客户的问题可以归纳为信息分散、方案割裂、试用决策成本高和持续维护责任不清。银伴的产品设计与商业模式分别对应这四类问题：把连续任务组合成路线、统一多种交互入口、提供低门槛验证包，并通过地图版本与责任人制度控制运维。",
    ],
    "3.3.1 现实问题": [
        "问题的关键不是用户完全不知道某个地点，而是无法把多个地点、楼层和流程节点组织成可执行顺序。尤其在从一楼大厅前往三楼科室后再返回一楼的任务中，如果系统省略电梯入口、目标楼层出口和返程步骤，单点位答案正确也可能导致整体路线失败。",
    ],
    "3.3.2 现有方案问题": [
        "现有方案各有适用边界，银伴并不以替代所有方案为目标。项目优势在于把自然语言问询、固定场景路线、适老界面和低成本实体导引结合起来，并在超出能力时主动转人工或屏幕提示，从而填补静态标识与高成本自主机器人之间的中间层。",
    ],
    "3.3.3 推广成本问题": [
        "推广成本主要发生在建立信任和跨部门协调，而不仅是宣传费用。首轮沟通材料需要同时提供功能清单、边界清单、验收指标、故障降级方式和预计维护投入，使门诊、信息、设备与安全相关人员能够基于同一范围评估项目。",
    ],
    "3.3.4 管理效率问题": [
        "管理效率应以单位时间内解决的有效任务和后续维护投入共同衡量。若设备减少了重复问询，却增加大量地图修改、充电或故障处理工作，整体效率未必提高。因此试点必须同步记录任务完成、人工接管、维护时长和异常类型。",
    ],
    "3.4 行业发展趋势": [
        "未来竞争重点将从单次设备销售转向场景数据维护、交互可达性、安全责任和服务连续性。对银伴而言，持续更新的地点知识、版本化楼层地图、可复用的验收流程和清晰的人工接管机制，比盲目增加摄像头或模型规模更具阶段价值。",
    ],
    "3.5 竞争分析": [
        "竞争比较采用自然语言能力、流程辅助、实体导引、部署成本和主要局限五个维度。银伴当前在适老问答和固定路线组合上具有差异化，但在陌生环境自主定位、动态避障和大规模部署经验方面仍弱于成熟通用机器人，后续验证应围绕这一真实边界展开。",
    ],
    "第四章 产品与技术分析": [
        "本章从产品构成、技术优势和可行性三个层面解释银伴如何工作。技术路线优先满足可解释、可离线、可停止和可验证四项要求，并将复杂自主导航、电梯联控和医疗信息系统接入留在取得真实场景授权后的后续阶段。",
    ],
    "4.1 产品概述": [
        "银伴由交互、服务、知识、规划、通信和执行六层组成。各层通过明确的数据与命令接口连接，使地点知识、路线规划和机器人控制可以独立测试；任一层发生异常时，系统能够定位问题并切换到较低风险的运行方式。",
    ],
    "4.1.1 产品介绍": [
        "产品的最小闭环是需求输入、意图识别、地点确认、路线生成、开始导引、状态更新和到达反馈。对跨楼层任务，系统在路线中显式加入电梯过渡，并在地图上切换楼层；对实体车未覆盖的部分，页面继续提供屏幕导引而不发送未经验证的运动命令。",
    ],
    "4.1.2 产品模块": [
        "模块化设计使比赛原型可以先验证最关键的服务逻辑。知识库内容变化不要求重写前端，路线节点变化不要求修改问答规则，通信方式变化也不影响用户输入；这种解耦降低联调风险，并为未来按医院配置地图与知识提供基础。",
    ],
    "4.1.3 使用说明": [
        "每一步都具有可见输入和可核验输出：用户可以查看系统识别到的原句，确认目的地顺序，阅读带楼层的路线节点，观察当前地图与机器人状态，并在任何阶段使用停止或复位按钮。该可见性有助于老年用户理解系统，也便于现场人员发现错误并接管。",
    ],
    "4.1.4 品牌与商标状态": [
        "品牌形象以亲和、稳定和方向感为设计原则。机器人外形降低陌生设备带来的距离感，定位标记与路线环直接对应导引功能；后续若进入商业使用，应完成商标近似检索、权属文件整理和素材授权归档，避免品牌宣传早于权利确认。",
    ],
    "4.1.5 产品标准与边界": [
        "内部标准服务于比赛验收和受控试点，不等同于医疗器械或公共场所移动设备的法定认证。每个对外版本都应附带适用场地、地图版本、最高速度、停止条件、不可用功能和责任联系人，使使用者能够判断产品能做什么以及何时不应继续使用。",
    ],
    "4.2 产品核心技术优势": [
        "银伴的技术创新属于面向场景约束的组合创新。系统没有追求在所有环境中自由导航，而是把确定性问答、分楼层图模型、人机交接、软硬件双层停车和多通道降级组合为可验证方案，以较低成本解决连续就医任务。",
    ],
    "4.2.1 本地确定性问答": [
        "地点名称、别名、楼层和流程答案以结构化数据保存，返回内容可以追溯到具体配置项。未知问题不通过自由生成补全，而是提示澄清或转人工；这使回答稳定性、修改责任和离线运行能力同时得到保障，尤其适合不能容忍路线幻觉的公共服务场景。",
    ],
    "4.2.2 多站点与分楼层路线": [
        "路线层使用带楼层属性的节点和边描述道路，电梯转换被建模为独立过渡而不是一条跨层直线。规划结果保留用户说出的目的地先后顺序，并允许同一路线经历一楼、三楼再返回一楼。该设计直接解决了传统单点回答无法表达连续任务的问题。",
    ],
    "4.2.3 软硬件闭环与安全状态机": [
        "上位网页负责呈现任务、发送高层命令和显示状态，ESP32负责电机、传感器和本地停止。两层都能触发停车，且STOP不依赖后续对话或网络返回；实体车只接受白名单路线编号，避免网页生成的任意路线直接转化为未经验证的运动。",
    ],
    "4.2.4 多通道降级与证据管理": [
        "多通道设计的目的不是增加功能数量，而是保证同一任务在部分能力失效时仍可完成。语音失败可改文字或按钮，蓝牙不稳可改USB，实体车未达标可改屏幕模拟。每次切换都需要保留原因、耗时和结果，形成可用于改进与答辩的证据链。",
    ],
    "4.3 产品可行性分析": [
        "可行性结论分为市场、生产和合规三条独立证据链。市场可行性回答用户是否需要、机构是否愿意试用；生产可行性回答样机能否稳定复现；合规可行性回答现阶段功能是否处于可控边界。三项同时成立后，项目才适合进入受控试点。",
    ],
    "4.3.1 市场可行性": [
        "访谈将采用任务复盘而不是直接询问是否喜欢机器人，以减少礼貌性偏差。验证重点包括用户是否能独立完成输入、是否理解路线、何时要求人工、愿意保持多远距离，以及机构愿意为哪些可量化结果付费。",
    ],
    "4.3.2 生产可行性": [
        "生产验证按照单模块、整车、路线和连续运行四级进行。只有电源、电机、巡线、测距和急停分别通过后才进入整车联调；参数标定一次只改变一个变量，并记录环境、速度、阈值和失败原因，以提高样机结果的可复现性。",
    ],
    "4.3.3 法律与合规可行性": [
        "合规策略采用最小数据和最小承诺原则。比赛版本不要求用户登录，不保存真实身份与病历，不连接医院业务系统；涉及真实场地、患者或内部地图时，必须先取得机构许可，并按用途限定、访问权限和保存期限管理资料。",
    ],
    "第五章 商业模式": [
        "本章将产品价值转换为可执行的机构采购路径。商业模式的核心不是立即销售大量设备，而是先用标准化验证包降低决策风险，再以明确的设备、部署和维护边界形成可核算、可验收和可续费的服务。",
    ],
    "5.1 商业模式概述": [
        "B2B2C模式对应银伴的实际价值链：机构承担采购、场地和维护责任，老年用户免费使用，团队提供设备、配置与支持。验证包不承诺替代导诊人员，而是围绕指定地点和路线测量任务完成、人工接管和维护投入，为后续采购提供依据。",
    ],
    "5.2 商业模式画布": [
        "画布中的九个模块必须形成闭环。目标客户决定产品边界，价值主张决定验收指标，渠道决定获客成本，关键活动与资源决定交付成本，客户关系和年度维护则决定是否能够形成持续收入。任一环节没有真实数据，都应保留为待验证假设。",
    ],
    "5.3 商业模式评价": [
        "当前模式的主要优点是固定范围、低首试成本和不依赖云服务；主要约束是机构决策周期、地图维护责任和单条实体路线的覆盖能力。项目只有在试点证明机构节省或改善的价值高于采购与维护成本后，才具备扩大交付的经济基础。",
    ],
    "5.4 合作伙伴类型": [
        "合作伙伴按作用分为场景方、验证方、技术供应方和推广方。医院或社区机构提供受控场景与流程信息，学校和指导教师帮助验证方法，供应商保障硬件一致性，赛事与校企平台承担早期触达；不同伙伴的责任和数据权限应分别书面确认。",
    ],
    "5.5 盈利方案": [
        "盈利能力依赖标准点位的单位经济模型，而不是依赖一次性高报价。团队需要逐单记录硬件采购、装配工时、地图配置、差旅培训、返修和远程支持成本；当直接成本长期高于8000元或验收返工频繁时，应先调整产品范围和流程，而不是用预测销量掩盖问题。",
    ],
    "5.6 发展计划": [
        "发展阶段采用门槛管理而非固定月份承诺。比赛原型通过后进入单楼层准备，取得场地书面同意后进入30天试用，形成安全与采购证据后再评估产品化。每一阶段都允许因证据不足而停止或回退，避免把时间表误写成必然结果。",
    ],
    "第六章 营销策略": [
        "本章的营销重点是建立机构信任，而不是追求公众曝光。所有材料围绕可验证任务、适用边界、验收标准和降级方案展开，使潜在合作方能够快速判断银伴是否适合其场景，并减少因能力误解造成的试点风险。",
    ],
    "6.1 主要销售策略": [
        "首轮销售漏斗分为目标机构清单、需求访谈、验证方案、书面试用意向和付费验证五个阶段。团队按阶段记录转化数量、未进入下一步的原因和决策周期，以区分产品问题、预算问题、合规问题和暂时没有需求，避免只统计联系数量。",
    ],
    "6.2 基于4Ps模型的营销策略": [
        "4Ps策略服务于同一目标：用明确产品范围降低使用风险，用分层价格降低决策风险，用定向渠道降低获客成本，用证据化宣传降低信任成本。四项策略必须与产品当前成熟度一致，不把比赛演示包装成真实商业部署。",
    ],
    "6.2.1 产品策略": [
        "标准化产品包明确一个楼层、地点数量、实体路线、问答范围、培训和验收内容。超出标准范围的新增楼层、路线或系统接口单独评估，防止定制需求侵占安全测试和维护资源；未完成的自动乘梯不进入报价或交付清单。",
    ],
    "6.2.2 价格策略": [
        "价格采用成本底线与客户价值双重校验。2980元验证包用于覆盖基础部署与测试，标准点位价格必须在直接成本和保修风险核算后确认；试点结束时根据任务完成、人工接管和维护数据复盘价格，不以竞赛阶段报价作为长期不变承诺。",
    ],
    "6.2.3 销售渠道": [
        "渠道优先级按照信任基础、场景匹配和决策可达性排序。校企合作与指导教师资源适合取得第一批访谈和受控演示机会，社区卫生与体检场景流程相对固定，适合早期验证；在形成案例之前，大范围投放难以提高有效转化。",
    ],
    "6.2.4 市场宣传": [
        "宣传内容采用问题场景、运行流程、量化证据和能力边界四部分结构。演示截图、测试结果和失败样本均标注时间与版本；涉及用户的材料必须匿名并取得使用许可，避免使用真实患者形象或可能引发医疗效果误解的表达。",
    ],
    "第七章 生产与服务": [
        "本章说明从参赛样机到受控交付所需的硬件核验、装配、配置、测试和售后责任。现阶段不建立生产线，重点是形成可以由团队重复执行的单台装配与单点位交付流程，并通过记录控制质量。",
    ],
    "7.1 产品生产经营计划": [
        "样机阶段采用通用模块和可拆卸固定方式，便于发现型号差异后快速调整。产品化前需要冻结主控、传感器、电源、底盘和充电方案，建立合格供应商与替代件清单，并把装配工时、返修原因和测试成本纳入单台成本。",
    ],
    "7.2 团队的生产技术能力": [
        "团队已经具备软件、协议和固件框架，但尚不能据此推断具备批量生产能力。当前需要补齐的是实物识别、焊接与线束规范、参数标定、连续运行和维修记录；若进入产品化阶段，还需引入结构、电气安全和质量管理方面的外部专业支持。",
    ],
    "7.3 品质控制和质量改进": [
        "质量控制遵循先安全、再稳定、后体验的顺序。急停、失联、障碍和丢线等问题未关闭前，不优化语音或动画效果；每次失败记录环境、版本、触发步骤、现象、恢复方式和责任人，同类问题连续出现时停止彩排并回到单模块验证。",
    ],
    "7.4 已购设备与到货核验": [
        "到货核验的目的不是确认数量即可，而是确认实物与设计假设一致。主控丝印、接口电平、电机额定电压、电池保护、传感器输出和充电器匹配关系必须逐项拍照记录；发现资料不一致时，以实测和供应商书面资料为准修改接线与固件。",
    ],
    "7.5 生产与交付流程": [
        "交付物不仅包括设备，还包括地图版本、知识库范围、接线与参数、测试结果、培训记录、故障降级方式和维护责任。试用期间的修改必须形成版本记录，避免现场配置与代码仓库不一致，从而保证问题能够复现和追责。",
    ],
    "第八章 财务分析": [
        "本章采用保守情景展示收入、直接成本、期间费用和资金安排。预测用于检验商业模型是否可能成立，不代表已经取得订单或承诺回报；后续每获得一项真实报价、试点或售后数据，都应更新假设并保留变动原因。",
    ],
    "8.1 财务基本情况与核算原则": [
        "财务记录按照权责来源和证据类型分类。已发生支出以订单、支付或发票为依据，尚未发生的产品化成本以供应商报价和明确假设估算，销售预测单独标记为经营情景。团队不得用成员无偿劳动掩盖长期交付成本，也不把赛事奖金计入主营收入。",
    ],
    "8.2 销售预计": [
        "销售预测采用验证包数量、验证转采购数量、标准点位数量和续费设备数四个驱动项。基准情景的重点是第一年完成付费验证并形成一个标准点位，而不是追求高增长；若转化率低于预期，应优先分析产品价值和采购障碍，再调整后续年度数量。",
    ],
    "8.3 成本费用核算": [
        "直接成本与期间费用分开核算。前者随验证包、设备或服务交付发生，后者用于研发、销售和管理；地图采集、差旅培训、返修和保修准备金不得遗漏。毛额率只能说明单次交付空间，不能替代现金流、回款周期和固定费用分析。",
    ],
    "8.4 利润及利润分配": [
        "利润分配以依法设立主体、完成纳税并保留经营安全垫为前提。研发与质量改进资金优先用于解决稳定性、合规和售后问题，团队激励建立在已实现并可分配的利润上；任何比例都需要在真实经营数据形成后由成员书面确认。",
    ],
    "8.5 融资方案和回报": [
        "融资服务于通过验证后的可靠性提升和受控试点，不用于填补尚未证明的市场需求。项目优先使用非股权资金，只有当付费意愿、单位成本、安全表现和复购或扩点意向同时形成证据后，才讨论股权融资与估值。",
    ],
    "8.5.1 资本结构与规模": [
        "学生团队阶段的投入应建立成员出资与项目资产台账，区分个人垫付、学校支持和赛事经费。若未来成立公司，应在知识产权归属、设备所有权和成员贡献确认后再设计股权结构，避免现在用缺乏依据的比例制造后续争议。",
    ],
    "8.5.2 融资形式": [
        "非股权资金与项目当前风险更匹配，因为其评价重点通常是研发里程碑和验证成果。申请材料应明确资金用途、交付节点和未达标处理；在商业模式尚未验证前引入股权，会放大估值分歧并降低团队后续调整空间。",
    ],
    "8.5.3 资金来源与使用计划": [
        "资金按里程碑分批使用：先完成结构和电子可靠性，再开展安全与合规准备，最后进入受控试点。每类支出设置预算上限、凭证要求和验收结果；若前一阶段未通过，不提前投入后续推广费用，以控制试错成本。",
    ],
    "8.5.4 投资收益与风险分析": [
        "当前阶段只能评估单位经济模型和主要风险，不能给出可信的收益率或回收期。潜在回报依赖试点转化、直接成本、回款周期、续费率和售后投入，其中任一变量变化都会显著影响结果，因此投资讨论必须建立在真实合同和连续运营数据上。",
    ],
    "第九章 实施计划与阶段结论": [
        "本章把到货后的技术任务转化为七天执行安排、量化验收、现场脚本、风险矩阵和明确阶段结论。计划的目标是确定比赛中哪些能力可以稳定展示，而不是在短时间内增加更多未经验证的功能。",
    ],
    "9.1 硬件到货后的七天计划": [
        "七天计划具有前置依赖关系：型号与供电未确认不能上电，单模块未通过不能整车联调，安全停车未通过不能进入连续彩排。每天结束时保存照片、参数、日志和失败记录，未完成项顺延并压缩展示范围，不通过并行冒险追赶进度。",
    ],
    "9.2 验证指标": [
        "指标分为功能正确性、运行稳定性、安全性和故障恢复四层。自动测试证明软件回归，预设问法和跨楼层截图证明交互与路线，20次彩排证明流程稳定，急停和异常测试证明风险可控；只有证据齐全的指标才能在答辩中表述为已完成。",
    ],
    "9.3 现场演示与降级预案": [
        "现场演示按照一个完整用户故事展开，不临时切换无关功能。主持人、操作员和安全观察员提前分工，操作员负责页面与小车，安全观察员始终能够触达急停；出现语音、通信或运动异常时按既定顺序降级，避免现场排错拖延主线。",
    ],
    "9.4 风险矩阵": [
        "风险处置优先级由影响而非发生概率单独决定。任何可能导致小车失控、碰撞或能力误解的风险均按高影响处理；每项风险在彩排前确认预防措施是否到位，触发后记录实际表现，并据此决定继续、降级或取消实体演示。",
    ],
    "9.5 阶段结论": [
        "因此，项目当前可以确认的结论是：软件网页端具备参赛演示条件，跨楼层路线逻辑已经完成并通过现有测试；实体车是否进入正式演示必须由到货后的安全和稳定性数据决定。在验收结果形成前，计划书、答辩和宣传统一使用这一口径。",
    ],
    "附录一 最小功能清单": [
        "最小功能清单同时作为范围冻结表。状态为软件已完成的功能需要保留版本与测试证据，状态为待硬件联调的功能必须在实车验收后才能勾选；任何新增想法先进入后续清单，不改变比赛主流程。",
    ],
    "附录二 软件目录": [
        "目录按照接口、核心规则、数据、服务、网页、固件和测试分层。各目录职责清晰，便于定位问答、路线、通信或执行问题；发布前只打包运行所需文件，并保留依赖版本与启动说明。",
    ],
    "附录五 答辩证据清单": [
        "证据按照版本、功能、安全、用户和来源五类归档，并确保每项材料能够对应计划书中的具体结论。截图和视频保留原始文件，测试输出保留完整日志，任何剪辑或摘要不得删除失败信息或改变结论。",
    ],
    "附录六 用户访谈提纲": [
        "访谈采用开放回忆、关键困难、现有应对和产品反馈的顺序，先理解真实任务，再展示方案。记录中同时保留支持、犹豫和拒绝意见，并区分受访者自己提出的问题与访谈者提示后的回答。",
    ],
    "附录七 两分钟演示操作清单": [
        "操作清单用于统一每次彩排条件。每次演示记录设备、软件版本、路线、耗时、是否人工干预和失败原因；出现异常时仍按同一脚本完成降级，以验证备用方案是否真正能在30秒内接管。",
    ],
}


def set_run_font(run, *, size=12, bold=False, color=BLACK, family=FONT_CN):
    run.font.name = family
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), family)
    rfonts.set(qn("w:ascii"), FONT_EN if family == FONT_CN else family)
    rfonts.set(qn("w:hAnsi"), FONT_EN if family == FONT_CN else family)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def clear_document_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def remove_header_footer_references(section):
    sect_pr = section._sectPr
    for tag in ("w:headerReference", "w:footerReference"):
        for node in list(sect_pr.findall(qn(tag))):
            sect_pr.remove(node)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5

    heading_settings = {
        "Heading 1": (22, WD_ALIGN_PARAGRAPH.CENTER, 17, 16.5),
        "Heading 2": (16, WD_ALIGN_PARAGRAPH.LEFT, 13, 8),
        "Heading 3": (14, WD_ALIGN_PARAGRAPH.LEFT, 10, 6),
    }
    for style_name, (size, alignment, before, after) in heading_settings.items():
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.4)


def add_page_number(footer):
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9)


def restart_page_numbering(section, start=1):
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is None:
        existing = OxmlElement("w:pgNumType")
        sect_pr.append(existing)
    existing.set(qn("w:start"), str(start))


def add_cover_line(doc, text, *, size, bold=False, before=0, after=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.5
    set_run_font(paragraph.add_run(text), size=size, bold=bold, family=FONT_COVER)
    return paragraph


def add_cover_field(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(2.2)
    paragraph.paragraph_format.right_indent = Cm(1.6)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5.2))
    set_run_font(paragraph.add_run(f"{label}："), size=14, bold=True, family=FONT_COVER)
    run = paragraph.add_run("\t" + value)
    set_run_font(run, size=13, family=FONT_CN)
    run.font.underline = True


def add_toc(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("目录"), size=22, bold=True, color=BLUE)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在Word中自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    set_run_font(run, size=11)


def add_heading(doc, text, level=1, *, new_page=False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = new_page
    set_run_font(
        paragraph.add_run(text),
        size={1: 22, 2: 16, 3: 14}.get(level, 12),
        bold=True,
    )
    return paragraph


def format_body_paragraph(paragraph, text, *, bold_lead=None, space_after=6):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), size=12, bold=True)
        text = text[len(bold_lead) :]
    set_run_font(paragraph.add_run(text), size=12)
    return paragraph


def add_para(doc, text, *, bold_lead=None, space_after=6):
    return format_body_paragraph(
        doc.add_paragraph(), text, bold_lead=bold_lead, space_after=space_after
    )


def insert_para_before(reference, text, *, bold_lead=None, space_after=6):
    return format_body_paragraph(
        reference.insert_paragraph_before(),
        text,
        bold_lead=bold_lead,
        space_after=space_after,
    )


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.8)
        paragraph.paragraph_format.first_line_indent = Cm(-0.42)
        paragraph.paragraph_format.line_spacing = 1.35
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        set_run_font(paragraph.add_run("• "), size=12, bold=True, color=BLUE)
        set_run_font(paragraph.add_run(item), size=12)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.9)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        paragraph.paragraph_format.line_spacing = 1.35
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        set_run_font(paragraph.add_run(f"{index}. "), size=12, bold=True, color=BLUE)
        set_run_font(paragraph.add_run(item), size=12)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)
        element.set(qn("w:space"), "0")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9.5, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(str(text)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for index, label in enumerate(headers):
        set_cell_text(header.cells[index], label, bold=True, color=WHITE, size=font_size, center=True)
        set_cell_shading(header.cells[index], BLUE)
        if widths:
            header.cells[index].width = widths[index]
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, value in enumerate(values):
            center = column_index == 0 and len(headers) > 2
            set_cell_text(row.cells[column_index], value, size=font_size, center=center)
            if widths:
                row.cells[column_index].width = widths[column_index]
            if row_index % 2 == 1:
                set_cell_shading(row.cells[column_index], LIGHT_GRAY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    set_run_font(paragraph.add_run(text), size=9)


def add_figure(doc, image_path, caption, *, alt_text, width=Inches(5.5)):
    image_path = Path(image_path)
    if not image_path.exists():
        raise FileNotFoundError(f"缺少网页成果截图：{image_path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    add_caption(doc, caption)


def heading_signature(doc):
    return [
        (paragraph.style.name, paragraph.text)
        for paragraph in doc.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]


def append_section_enrichments(doc):
    all_paragraphs = list(doc.paragraphs)
    heading_rows = [
        (index, paragraph)
        for index, paragraph in enumerate(all_paragraphs)
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    headings = [paragraph for _, paragraph in heading_rows]
    available = {paragraph.text for paragraph in headings}
    missing = sorted(set(SECTION_ENRICHMENTS) - available)
    if missing:
        raise RuntimeError(f"扩充内容对应不到原目录标题：{missing}")

    for row_index, (index, heading) in enumerate(heading_rows):
        texts = SECTION_ENRICHMENTS.get(heading.text, [])
        if not texts:
            continue
        reference = heading_rows[row_index + 1][1] if row_index + 1 < len(heading_rows) else None
        for text in texts:
            if reference is None:
                add_para(doc, text)
            else:
                insert_para_before(reference, text)


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, index, title, organization, date, url):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.first_line_indent = Cm(-0.8)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    set_run_font(paragraph.add_run(f"[{index}] {organization}：《{title}》，{date}。 "), size=9.5)
    add_hyperlink(paragraph, url, url)


def build_document():
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(TEMPLATE_PATH)
    digest = sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
    if digest != EXPECTED_TEMPLATE_SHA256:
        raise RuntimeError("模板文件发生变化，请重新执行模板蒸馏。")

    doc = Document(TEMPLATE_PATH)
    clear_document_body(doc)
    configure_styles(doc)
    configure_page(doc.sections[0])
    remove_header_footer_references(doc.sections[0])

    for _ in range(2):
        doc.add_paragraph()
    add_cover_line(doc, "中国国际大学生创新大赛（2026）", size=22, before=10)
    add_cover_line(doc, "参赛项目", size=22, after=32)
    add_cover_line(doc, "商业计划书", size=22, after=34)
    add_cover_field(doc, "学院", "待填写")
    add_cover_field(doc, "项目名称", "银伴——医院智能陪诊导引机器人")
    add_cover_field(doc, "所在赛道", "产业赛道")
    add_cover_field(doc, "项目组别", "成果转化组")
    add_cover_field(doc, "项目类别", "待学院确认")
    add_cover_field(doc, "项目负责人", "翟倩妮")
    add_cover_field(doc, "联系电话", "待填写")
    add_cover_line(doc, "兰州理工大学", size=16, before=26, after=3)
    add_cover_line(doc, "2026年9月", size=16)

    doc.add_page_break()
    add_toc(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body_section)
    restart_page_numbering(body_section, 1)
    add_page_number(body_section.footer)

    # 第一章
    add_heading(doc, "第一章 执行概要", 1)
    add_heading(doc, "1.1 项目概况", 2)
    add_para(
        doc,
        "银伴是一套面向老年就医者的医院智能陪诊导引原型。随着数字技术快速进入挂号、报到、缴费、检查和取药流程，一部分老年人难以及时适应不断变化的终端和操作方式；同时，医院楼栋、楼层和科室关系复杂，单纯增加自助设备并不能自动解决理解路线与执行下一步的问题。银伴用语音、文字和大按钮接收需求，结合本地问答、分楼层路线图和低速实体小车，帮助用户完成问地点、看路线、跟随导引和随时停止。",
    )
    add_para(
        doc,
        "项目当前定位为固定环境参赛原型，而不是可直接在医院自主运行的成熟机器人。软件端已完成并通过28项自动测试；硬件已采购，仍需在到货后完成型号核验、接线、巡线、避障和连续运行验证。",
    )

    add_heading(doc, "1.2 产品介绍", 2)
    add_para(
        doc,
        "产品由浏览器交互端、本地服务端、医院知识与路线数据、机器人通信层和ESP32执行端组成。屏幕端能够按用户说话顺序解析多个目的地，并在跨楼层时明确展示进入电梯、到达目标楼层和返回一楼等过渡；实体车只执行已验证的同楼层固定路线，到达电梯入口后停止并由屏幕继续提示，不宣称自动乘梯。",
    )

    add_heading(doc, "1.3 市场分析", 2)
    add_para(
        doc,
        "首批使用者是首次到院、需要跨多个地点行动或不熟悉智能终端的老年患者及陪同家属；首批付费与决策主体是医院门诊部、社区卫生服务中心和体检中心。项目不使用全国诊疗人次直接推算销售额，而以单点位30天验证、用户任务完成率、人工接管率和机构试用意向验证真实需求。",
    )

    add_heading(doc, "1.4 商业模式", 2)
    add_para(
        doc,
        "银伴采用机构付费、老年用户免费使用的模式。首轮执行价格为：30天验证包2980元每点位、标准设备12800元每台、场景部署3000元每点位、第二年起年度软件服务1200元每台。一个标准点位第一年合同额为15800元；在产品化直接成本控制在8000元以内时，单点位目标贡献毛额为7800元。上述价格属于验证期报价，不代表已经形成销售收入。",
    )

    add_heading(doc, "1.5 组织管理", 2)
    add_para(
        doc,
        "团队由翟倩妮、王子阳、王钰晋三人组成。翟倩妮负责需求、范围、对外沟通、商业验证和答辩；王子阳负责软件、数据、接口和版本管理；王钰晋负责硬件、固件、安全测试和实车联调。材料制作与现场演示由三人交叉复核，任何带电或运动测试至少两人在场。",
    )

    add_heading(doc, "1.6 风险管理", 2)
    add_para(
        doc,
        "核心风险是实物型号不符、巡线不稳定、语音受噪声影响、通信中断、地图过期和能力被误解。项目采用STOP最高优先级、障碍与丢线本地停车、USB与模拟模式降级、地图版本记录和明确非诊疗边界控制风险。实体演示只有在20次完整运行至少成功18次且全部安全测试通过后才能进入比赛主流程。",
    )

    # 第二章
    add_heading(doc, "第二章 项目概况", 1, new_page=True)
    add_heading(doc, "2.1 项目介绍", 2)
    add_para(
        doc,
        "银伴围绕复杂公共空间中的适老信息辅助与行动导引展开。医院版本先聚焦五类能力：问地点、办流程、带路线、陪等待、接求助。当前原型重点验证前三类能力，并为停止与人工求助保留固定入口。项目核心不在于堆叠摄像头或大模型，而在于把可解释的本地问答、路线数据、适老交互和安全执行形成稳定闭环。",
    )
    add_para(
        doc,
        "项目的竞争优势来自对真实就医任务的拆解与重组。银伴不只回答某个科室位于几楼，而是保留用户提出的多个目的地及先后顺序，把大厅、通道、电梯入口、目标楼层出口和科室连接成连续步骤；同时通过大字、大按钮、语音与文字并行入口降低操作门槛。其创新重点属于场景化系统创新，即用确定性知识、本地路线、安全状态机和低成本实体载体共同解决老年用户在复杂空间中“听得懂但走不到”的问题。",
    )
    add_table(
        doc,
        ["优势与创新维度", "具体实现", "相较常见方案的改进", "当前证据"],
        [
            ("适老交互", "语音、文字、大按钮进入同一任务链；界面采用大字号和高对比度", "减少菜单层级和新应用学习成本", "网页实际运行与输入输出截图"),
            ("多站连续任务", "识别多个目的地并保持用户表达顺序", "由单点问路扩展为可执行的连续就医路线", "自动测试与多目的地页面结果"),
            ("跨楼层可解释路线", "将电梯入口、乘梯方向、楼层出口和返程建模为独立节点", "避免用跨楼层直线连接造成路线误导", "一楼、乘梯、三楼分阶段截图"),
            ("本地确定性服务", "地点、别名、流程和楼层信息来自可审核JSON配置", "断网可用，答案可追溯，降低路线幻觉风险", "知识库、接口与28项测试"),
            ("软硬件双层安全", "网页提供启动、停止和复位；ESP32保留本地STOP与异常停车", "上位机失联时仍能停车，责任边界更清楚", "协议已完成，实车功能待到货验收"),
            ("低成本可复制原型", "采用浏览器、本地服务与通用ESP32底盘，硬件预算不超过500元", "用较低成本完成比赛级闭环验证", "采购清单、代码仓库与验收表"),
        ],
        widths=[Cm(3.1), Cm(5.3), Cm(5.0), Cm(3.2)],
        font_size=7.8,
    )

    add_heading(doc, "2.2 当前成果", 2)
    add_table(
        doc,
        ["成果类别", "当前状态", "可核验依据"],
        [
            ("软件", "已完成", "适老网页、对话、分楼层路线、模拟机器人和API可运行"),
            ("测试", "已完成", "28项自动测试全部通过；JavaScript语法检查通过"),
            ("路线", "已完成", "多目的地顺序、跨楼层电梯过渡和楼层切换已验证"),
            ("品牌", "已完成", "银伴机器人图标、方形版、横向字标和封面形象"),
            ("采购", "已完成", "硬件在运输途中，原型总预算上限500元"),
            ("实物联调", "待完成", "需按实物核对引脚、电平、电机、传感器和电源"),
        ],
        widths=[Cm(3.0), Cm(3.0), Cm(10.5)],
    )
    add_para(
        doc,
        "网页端当前已经形成完整的软件演示流程。启动本地服务后，用户在浏览器中通过语音或文字输入需求；后端识别全部目的地并保持原有顺序，随后返回地点说明、楼层信息和完整路线。用户点击开始导引后，页面进入运行状态，定位图标沿当前楼层的道路节点移动；到达电梯节点时，页面显示正在乘梯及楼层变化，完成过渡后自动切换到目标楼层地图。用户可随时停止、继续或复位，实体车未接入时由模拟机器人复现同一状态流程。",
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-01-input-answer.png",
        "图1 网页端识别多目的地需求并返回完整地点与楼层信息",
        alt_text="银伴网页端输入先去心内科再去卫生间后，系统同时返回两个目的地和跨楼层提示",
        width=Inches(5.75),
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-02-map-floor1-mobile.png",
        "图2 一楼路线图按真实道路节点连接门诊大厅与三号电梯口",
        alt_text="银伴一楼路线图，显示门诊大厅、三号电梯口及沿道路节点绘制的路线",
        width=Inches(5.55),
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-04-guiding-controls-mobile.png",
        "图3 点击开始导引后页面进入运行状态并提供停止与复位控制",
        alt_text="银伴网页端正在导引状态，显示运行信息、停止和复位按钮",
        width=Inches(5.55),
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-05-guiding-map-mobile.png",
        "图4 导引过程中机器人图标沿一楼路线节点移动",
        alt_text="银伴网页端导引地图，机器人定位图标正在一楼道路节点上移动",
        width=Inches(5.55),
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-06-elevator-transition.png",
        "图5 到达电梯节点后显示由一楼前往三楼的乘梯过渡",
        alt_text="银伴网页端电梯过渡提示，显示正在乘坐三号电梯由一楼前往三楼",
        width=Inches(5.55),
    )
    add_figure(
        doc,
        SCREENSHOT_DIR / "yinban-web-03-map-floor3-mobile.png",
        "图6 完成楼层过渡后自动切换至三楼心内科路线图",
        alt_text="银伴三楼路线图，显示三号电梯出口与心内科之间的道路节点和路线",
        width=Inches(5.55),
    )
    add_para(
        doc,
        "以上图片来自2026年9月10日当前代码版本的实际运行界面，不是概念效果图。它们能够证明软件端已经实现多目的地解析、分楼层路线、开始导引、图标移动和电梯过渡；但不能据此推定实体车已经完成巡线、避障或连续运行验收，实物能力仍需按第九章所列指标验证。",
    )

    add_heading(doc, "2.3 组织架构", 2, new_page=True)
    add_table(
        doc,
        ["成员", "主责岗位", "近期交付", "复核责任"],
        [
            ("翟倩妮", "项目负责人、产品与商务", "范围冻结、计划书、用户访谈、机构沟通、答辩", "复核测试证据和价格口径"),
            ("王子阳", "软件与系统联调", "前后端、路线数据、接口、自动测试和发布包", "复核演示材料中的技术表述"),
            ("王钰晋", "硬件、固件与安全", "接线、巡线、避障、急停、日志和实车视频", "复核现场安全和硬件参数"),
        ],
        widths=[Cm(2.6), Cm(4.0), Cm(6.0), Cm(4.0)],
        font_size=8.8,
    )
    add_para(
        doc,
        "三人均参与最终彩排。硬件负责人不得在带电与运动测试中同时担任唯一安全观察者；软件负责人冻结版本后只修复影响主流程或安全的问题；项目负责人确保计划书中的每项已完成结论均能回到代码、日志、照片或视频证据。",
    )

    add_heading(doc, "2.4 项目宗旨", 2, new_page=True)
    add_para(
        doc,
        "项目坚持技术辅助而不替代人工、提供流程信息而不进行医疗诊断、先验证固定路线再讨论复杂自主导航。银伴希望让老年人无需掌握新的复杂系统，也能通过说话、按钮或文字获得清楚的下一步指引。",
    )

    # 第三章
    add_heading(doc, "第三章 市场分析", 1, new_page=True)
    add_heading(doc, "3.1 行业背景概述", 2)
    add_para(
        doc,
        "国家统计局发布的2025年国民经济和社会发展统计公报显示，2025年末我国60岁及以上人口为32338万人，占总人口23.0%；65岁及以上人口为22365万人，占15.9%；全年医疗卫生机构总诊疗人次为105.8亿。[1] 这些数据说明医院适老服务面对长期且高频的使用场景，但不能直接等同为机器人市场规模。",
    )
    add_para(
        doc,
        "数字化发展速度快于部分老年用户学习和适应的速度。医院不断引入线上预约、自助报到、电子凭证和移动支付，流程效率提高的同时，也可能把操作负担转移给不熟悉智能设备的人。银伴不是要求用户再学习一个复杂应用，而是把流程、地点和路线转换为直观的语音、文字和行动提示。",
    )

    add_heading(doc, "3.2 目标客户", 2)
    add_table(
        doc,
        ["对象", "核心任务", "采购或使用关注点"],
        [
            ("老年就医者", "问地点、问流程、跟随导引、停止或求助", "大字、少步骤、慢语速、安全和人工兜底"),
            ("陪同家属", "安排多个地点并确认先后顺序", "路线一致、临时调整和清晰提示"),
            ("导诊与志愿者", "处理高频重复问询", "知识库可维护、异常可接管"),
            ("医院门诊部", "改善服务体验并减轻重复指路", "安全、运维、地图责任和合规"),
            ("社区与体检中心", "在人员有限场景提供固定流程辅助", "成本、易用性和本地维护"),
        ],
        widths=[Cm(3.2), Cm(6.5), Cm(6.8)],
        font_size=8.8,
    )

    add_heading(doc, "3.3 目标客户面临的问题", 2)
    add_heading(doc, "3.3.1 现实问题", 3)
    add_para(
        doc,
        "用户在到院、挂号、报到、候诊、检查、缴费和取药之间频繁切换，入口、楼栋、楼层、凭证和顺序信息分散。焦虑、视听能力下降或首次到院会进一步增加操作负担。单次回答科室位置不足以解决连续任务，路线必须表达进入电梯、上下楼和返回大厅等过渡。",
    )

    add_heading(doc, "3.3.2 现有方案问题", 3)
    add_para(
        doc,
        "人工导诊自然但在高峰期服务能力有限；静态标识成本低但不能根据目的地组合路线；自助机信息集中却需要用户理解菜单；手机地图依赖个人设备和室内定位；通用自主导航机器人功能较强，但设备、部署、地图维护和安全成本更高。银伴先以固定路线验证需求，不把昂贵能力作为参赛前提。",
    )

    add_heading(doc, "3.3.3 推广成本问题", 3)
    add_para(
        doc,
        "医疗机构引入新设备需要经过门诊、信息、设备、安全和后勤等多方评估，直接销售周期长。项目用2980元、30天、单点位的验证包降低首次决策成本，并以明确验收表替代抽象功能介绍。首轮获客依靠校企对接、指导教师资源和机构定向访谈，不以大规模广告作为主要渠道。",
    )

    add_heading(doc, "3.3.4 管理效率问题", 3)
    add_para(
        doc,
        "地图变化、科室调整、设备充电、故障处理和日志留存都会产生持续运维成本。每个试点地图必须标明来源、采集日期、适用楼层、审核人和更新责任人；道路或科室变化后先停用旧版本，再复核路线。只有当机构能够承担这一维护机制时，设备才适合继续部署。",
    )

    add_heading(doc, "3.4 行业发展趋势", 2)
    add_para(
        doc,
        "适老化服务的方向不是取消数字化，而是提供多入口、可理解和可转人工的数字服务。无障碍环境建设相关法律要求公共服务场所的自助终端具备语音、大字等无障碍功能，并保留现场指导和人工办理方式。[4] 因此，银伴的中期价值将更多来自适老交互、场景配置和运维数据，而不是单一硬件外形。",
    )

    add_heading(doc, "3.5 竞争分析", 2)
    add_table(
        doc,
        ["方案", "自然语言", "流程辅助", "实体导引", "成本与部署", "主要局限"],
        [
            ("人工导诊", "强", "强", "有限", "持续人员成本", "高峰期服务能力受限"),
            ("自助机", "弱", "中", "无", "已有基础设施", "菜单学习成本高"),
            ("手机地图", "中", "弱", "无", "边际成本低", "依赖手机与室内定位"),
            ("通用自主机器人", "强", "中", "强", "设备和部署成本高", "地图运维和场地要求高"),
            ("银伴原型", "强", "强", "固定路线", "硬件500元内", "当前仅适合受控场地验证"),
        ],
        widths=[Cm(2.5), Cm(2.0), Cm(2.0), Cm(2.1), Cm(3.3), Cm(4.5)],
        font_size=8.0,
    )

    # 第四章
    add_heading(doc, "第四章 产品与技术分析", 1, new_page=True)
    add_heading(doc, "4.1 产品概述", 2)
    add_heading(doc, "4.1.1 产品介绍", 3)
    add_para(
        doc,
        "银伴提供医院常见地点与流程问答、按顺序规划多个目的地、分楼层地图显示、屏幕模拟导引、实体车固定路线导引、障碍停车和人工停止。核心问询由本地知识库与规则完成，断网时仍可运行。",
    )

    add_heading(doc, "4.1.2 产品模块", 3)
    add_table(
        doc,
        ["层级", "组件", "职责", "降级方式"],
        [
            ("交互层", "浏览器界面与语音接口", "接收语音、文字和按钮并展示路线状态", "语音失败时使用文字或按钮"),
            ("服务层", "FastAPI与Pydantic", "组织问答、路线和机器人命令", "本地运行，不依赖外网"),
            ("知识层", "医院地点与流程JSON", "提供可审核事实、别名和楼层", "未知问题转澄清或人工"),
            ("规划层", "图节点与楼层过渡", "组合多站路线并生成电梯转换步骤", "实体模式到电梯口交接"),
            ("通信层", "USB串口或蓝牙SPP", "发送命令并读取状态", "蓝牙失败改USB或模拟"),
            ("执行层", "ESP32状态机", "巡线、避障、停车和到达反馈", "本地STOP与超时独立生效"),
        ],
        widths=[Cm(2.3), Cm(4.0), Cm(6.0), Cm(4.3)],
        font_size=8.3,
    )

    add_heading(doc, "4.1.3 使用说明", 3)
    add_numbered(
        doc,
        [
            "用户说出或输入一个或多个目的地，例如先去心内科再去卫生间。",
            "系统识别目的地及顺序，同时优先检查停止和紧急求助关键词。",
            "路线服务根据节点、道路与楼层连接返回完整步骤，并明确电梯上下楼过渡。",
            "用户点击开始导引；屏幕图标沿真实路线节点移动，跨楼层时切换楼层地图。",
            "实体模式仅执行已验证的同楼层路线；到达电梯入口后停车并转为屏幕续引。",
            "遇障碍、丢线、通信中断、超时或人工STOP时立即停车。",
        ],
    )

    add_heading(doc, "4.1.4 品牌与商标状态", 3)
    add_para(
        doc,
        "银伴文字、图形标志和机器人形象由项目团队形成并纳入仓库管理。目前尚未提交商标注册申请，计划在对外商业使用前完成近似检索与类别确认。计划书不使用未取得的注册商标标识，也不声称已经获得专利或软件著作权。",
    )

    add_heading(doc, "4.1.5 产品标准与边界", 3)
    add_para(
        doc,
        "参赛原型采用内部验收标准：核心问法全部正确、指定路线20次至少成功18次、急停与四类异常停车全部有效、完整演示不超过120秒、故障30秒内完成降级。真实销售前需另行完成产品电气安全、充电、电磁兼容、场地运行责任和数据合规评估；当前测试不能替代法定认证。",
    )

    add_heading(doc, "4.2 产品核心技术优势", 2)
    add_heading(doc, "4.2.1 本地确定性问答", 3)
    add_para(doc, "高频地点和流程优先由本地规则与可审核JSON知识库回答，避免云端模型产生不可预测的路线或诊疗内容。安全关键词优先级高于普通导航。")
    add_heading(doc, "4.2.2 多站点与分楼层路线", 3)
    add_para(doc, "路线规划依据显式节点和边组合路径，保持用户表达的目的地顺序。跨楼层路线包含电梯入口、乘梯方向、目标楼层出口和返程步骤，地图按楼层切换，不再用跨楼层直线连接造成误导。")
    add_heading(doc, "4.2.3 软硬件闭环与安全状态机", 3)
    add_para(doc, "网页通过统一协议启动、停止、继续和复位机器人；ESP32独立判断障碍、丢线、通信中断和超时，在上位机失联时仍能停车。状态回传使页面与实体行为保持一致。")
    add_heading(doc, "4.2.4 多通道降级与证据管理", 3)
    add_para(doc, "语音、文字和按钮可以完成同一任务；蓝牙、USB和模拟机器人形成三级执行路径。测试日志、截图、视频、接线照片和Git提交用于证明每项结论，预测值不写成实测结果。")

    add_heading(doc, "4.3 产品可行性分析", 2)
    add_heading(doc, "4.3.1 市场可行性", 3)
    add_para(doc, "需求成立与否将由5至8名老年用户或家属访谈、1至2名服务人员访谈以及至少3名机构决策人员沟通验证。首轮判断指标是任务完成率、求助次数、跟随接受条件和机构是否愿意进入30天受控试用，而不是礼貌性好评。")
    add_heading(doc, "4.3.2 生产可行性", 3)
    add_para(doc, "比赛样机使用通用ESP32、两轮底盘、巡线模块、超声波传感器和低压供电，预算上限500元，器件可替换且维修门槛低。产品化版本直接硬件、装配、地图部署、差旅和保修准备金合计目标不超过8000元。")
    add_heading(doc, "4.3.3 法律与合规可行性", 3)
    add_para(doc, "系统只提供地点与流程信息，不诊断、不处方、不采集身份、医保、病历、人脸或真实预约数据。真实地图须使用公开资料并记录来源，或取得书面授权；电梯联动、医院信息系统接口和真实患者数据接入不属于首轮产品承诺。")

    # 第五章
    add_heading(doc, "第五章 商业模式", 1, new_page=True)
    add_heading(doc, "5.1 商业模式概述", 2)
    add_para(
        doc,
        "银伴采用机构采购或租用、老年用户免费使用的B2B2C模式。团队先交付单楼层、固定路线、非诊疗、非联网的30天验证包；机构确认安全和使用价值后，再采购标准设备、场景部署和年度软件服务。",
    )
    add_table(
        doc,
        ["收费项目", "执行价格", "包含内容", "收费与验收规则"],
        [
            ("30天验证包", "2980元/点位", "1台设备、1个楼层、10个以内地点、1条实体路线、培训和验证报告", "部署完成收费；90天内采购可抵扣1500元"),
            ("标准设备", "12800元/台", "受控区域移动底盘、交互终端、充电装置、1年基础保修", "签约50%、部署40%、30天验收10%"),
            ("场景部署", "3000元/点位", "首层地图、10个地点、1条实体路线、知识库初始化和培训", "新增楼层1500元；新增实体路线800元"),
            ("年度软件服务", "1200元/台/年", "版本更新、知识库维护、日志导出与远程支持", "首年包含；第二年起按年续费"),
        ],
        widths=[Cm(2.8), Cm(2.8), Cm(6.5), Cm(4.5)],
        font_size=8.0,
    )

    add_heading(doc, "5.2 商业模式画布", 2)
    add_table(
        doc,
        ["模块", "银伴的具体设定"],
        [
            ("客户细分", "医院门诊部、社区卫生服务中心、体检中心；使用者为老年患者和家属"),
            ("价值主张", "低成本、可离线、适老、多站点、分楼层屏幕导航与固定路线实体导引"),
            ("渠道", "校企对接、指导教师资源、医院志愿服务与门诊部门定向沟通"),
            ("客户关系", "30天受控验证、培训、验收报告、年度维护"),
            ("收入来源", "验证包、设备采购、场景部署、年度软件服务"),
            ("核心资源", "适老交互组件、医院知识与地图配置、机器人协议、安全测试证据"),
            ("关键活动", "需求访谈、地图配置、软硬件联调、现场验收和版本维护"),
            ("关键伙伴", "学校、指导教师、医院或社区服务机构、硬件供应商"),
            ("成本结构", "硬件与装配、地图部署、差旅培训、软件维护、保修与安全测试"),
        ],
        widths=[Cm(3.6), Cm(13.0)],
        font_size=8.8,
    )

    add_heading(doc, "5.3 商业模式评价", 2)
    add_para(
        doc,
        "模式优点是首轮验证成本可控、核心功能不依赖云服务、机构可以从单点位开始；主要限制是医疗机构决策周期长、地图与设备需要持续维护，且固定路线能覆盖的任务有限。项目只有在取得至少1份书面试用意向、完成受控用户验证并核算真实交付成本后，才进入产品化评估。",
    )

    add_heading(doc, "5.4 合作伙伴类型", 2)
    add_para(
        doc,
        "当前没有已签约合作伙伴。优先接触对象包括兰州地区医院门诊或志愿服务部门、社区卫生服务中心、体检中心、适老服务研究团队和可靠的ESP32硬件供应商。对外材料必须区分已接触、已表达兴趣、书面试用意向和正式合同四种状态。",
    )

    add_heading(doc, "5.5 盈利方案", 2)
    add_para(
        doc,
        "一个标准点位第一年合同额为15800元。目标直接成本上限为8000元，其中产品化硬件与装配不超过5500元、地图部署和差旅不超过1500元、首年保修准备金不超过1000元；目标贡献毛额7800元，约占合同额49%。该比例未扣除研发、销售和管理费用。",
    )
    add_para(
        doc,
        "标准验收条件为：预设地点问答15条全部正确；指定路线20次至少成功18次；急停、障碍、丢线、超时和断连全部触发停车；部署清单、培训记录和故障说明完成交接。未达到验收条件不得确认尾款。",
    )

    add_heading(doc, "5.6 发展计划", 2)
    add_table(
        doc,
        ["阶段", "交付重点", "进入下一阶段的条件"],
        [
            ("比赛原型", "软件闭环、实车单路线、安全测试和答辩证据", "实体门槛全部通过或按规则降级为软件演示"),
            ("单楼层准备", "公开或授权地图、运维责任、用户与工作人员访谈", "场地方书面同意受控测试"),
            ("30天受控试用", "匿名任务数据、人工接管、设备与地图维护记录", "安全无重大问题且形成继续试用或采购意向"),
            ("产品化评估", "结构加固、多设备管理、认证与售后成本", "存在明确采购主体且单位经济性成立"),
        ],
        widths=[Cm(3.0), Cm(7.0), Cm(6.6)],
        font_size=8.6,
    )

    # 第六章
    add_heading(doc, "第六章 营销策略", 1, new_page=True)
    add_heading(doc, "6.1 主要销售策略", 2)
    add_para(
        doc,
        "首轮销售不以广告曝光为目标，而以完成机构问题访谈和取得受控试用机会为目标。团队建立8家目标机构名单，完成至少3名决策或使用部门人员访谈，向至少3家发送统一验证包方案，并争取1份非约束性书面试用意向。每次沟通记录反对意见、预算范围、决策链和下一步。",
    )
    add_table(
        doc,
        ["任务", "数量目标", "完成证据", "负责人"],
        [
            ("建立机构名单", "8家", "单位、部门、联系日期和下一步", "翟倩妮"),
            ("决策人员访谈", "至少3人", "匿名纪要、采购关注点和反对意见", "翟倩妮"),
            ("发送验证方案", "至少3份", "统一报价、边界和30天清单", "翟倩妮、王子阳"),
            ("书面试用意向", "至少1份", "邮件、回函或非约束意向说明", "翟倩妮"),
            ("价格复盘", "1份", "客户反馈、实际成本和调整依据", "全体成员"),
        ],
        widths=[Cm(3.5), Cm(2.5), Cm(7.2), Cm(3.4)],
        font_size=8.4,
    )

    add_heading(doc, "6.2 基于4Ps模型的营销策略", 2)
    add_heading(doc, "6.2.1 产品策略", 3)
    add_para(doc, "对外只销售通过验收的单点位版本，标准范围固定为一个楼层、10个以内地点和一条实体路线。跨楼层先通过屏幕提示与电梯口交接完成，不把自动乘梯作为销售功能。")
    add_heading(doc, "6.2.2 价格策略", 3)
    add_para(doc, "用2980元验证包降低第一次尝试成本，用12800元设备和3000元部署构成标准合同。新增楼层、路线和地点按统一规则计价，避免每次重新议价导致范围失控。")
    add_heading(doc, "6.2.3 销售渠道", 3)
    add_para(doc, "优先使用校企命题对接、创新创业赛事、指导教师与校友资源、医院志愿服务部门和社区卫生服务机构的定向沟通。未取得机构许可前，不进入真实门诊环境采集数据或开展宣传演示。")
    add_heading(doc, "6.2.4 市场宣传", 3)
    add_para(doc, "宣传材料采用两分钟演示视频、单页验证包说明、技术边界清单和匿名测试结果。所有画面注明模拟环境或受控试点，不使用真实患者形象，不宣称已在医院部署，也不把计划指标描述为完成数据。")

    # 第七章
    add_heading(doc, "第七章 生产与服务", 1, new_page=True)
    add_heading(doc, "7.1 产品生产经营计划", 2)
    add_para(
        doc,
        "比赛阶段采用小批量手工装配与本地部署：先核验主控、电机、驱动、巡线、测距和电源，再完成单模块测试、整车联调、地图配置、20次彩排和证据归档。产品化前不建立生产线；若形成订单，底盘与标准电子模块外购，团队负责软件镜像、线路检查、场景配置和验收。",
    )

    add_heading(doc, "7.2 团队的生产技术能力", 2)
    add_para(
        doc,
        "团队已经完成软件架构、API、分楼层路线、模拟机器人、通信协议和ESP32固件骨架，28项自动测试通过。尚未完成的能力是实物型号确认、电气接线、巡线参数标定、障碍阈值和长时间运行。团队以到货后的七天计划补齐这些环节，不把代码完成等同于实车完成。",
    )

    add_heading(doc, "7.3 品质控制和质量改进", 2)
    add_table(
        doc,
        ["质量门槛", "目标", "记录方式", "不通过处理"],
        [
            ("软件回归", "28项及新增测试全部通过", "pytest输出和提交号", "禁止发布演示包"),
            ("问答路线", "15条预设问法全部正确", "输入输出与截图", "修正知识或路线数据"),
            ("完整运行", "20次至少18次成功", "逐次彩排表", "缩短路线或改用模拟"),
            ("安全停车", "急停、障碍、丢线、超时、断连全部有效", "日志和视频", "实体车退出主流程"),
            ("故障降级", "30秒内切换USB或模拟", "计时记录", "提前改为软件主演示"),
        ],
        widths=[Cm(3.0), Cm(4.0), Cm(5.0), Cm(4.6)],
        font_size=8.5,
    )

    add_heading(doc, "7.4 已购设备与到货核验", 2)
    add_table(
        doc,
        ["模块", "目标规格", "预算上限", "到货核验"],
        [
            ("主套件", "ESP32 WROOM两轮车、五路巡线、超声波", "260元", "主控丝印、焊接、资料和电池方案"),
            ("供电与安全", "保护电池、匹配充电器、电源开关、急停", "80元", "极性、电压、保护与开关"),
            ("传感器与备件", "3.3V兼容巡线、测距、按键、线材", "70元", "输出电平、阈值、接口和数量"),
            ("路线与展台", "胶带、KT板、固定件、标识", "60元", "对比度、转弯半径和固定强度"),
            ("应急额度", "临时替换接头、线材或固定件", "30元", "仅用于故障替换并记录原因"),
        ],
        widths=[Cm(3.0), Cm(5.8), Cm(2.6), Cm(5.2)],
        font_size=8.3,
    )

    add_heading(doc, "7.5 生产与交付流程", 2)
    add_numbered(
        doc,
        [
            "按订单或验证范围冻结地点、楼层、实体路线和不包含功能。",
            "核验硬件型号、电平、电源和接线，完成单模块安全测试。",
            "导入场景知识与地图，逐条验证问答、楼层过渡和路线节点。",
            "完成20次运行、安全停车和故障降级测试，形成验收报告。",
            "交付设备、地图版本、培训记录、故障说明和维护责任表。",
            "试用期记录匿名任务结果；地图变化时暂停旧路线并重新审核。",
        ],
    )

    # 第八章
    add_heading(doc, "第八章 财务分析", 1, new_page=True)
    add_heading(doc, "8.1 财务基本情况与核算原则", 2)
    add_para(
        doc,
        "项目目前处于学生团队原型阶段，尚未设立经营公司、没有历史营业收入、贷款或股权融资。参赛样机预算上限为500元，订单、发票或付款记录按硬件、供电安全、传感器、展台和应急五类登记。以下三年数据是按当前报价形成的保守情景预测，不是已实现业绩。",
    )

    add_heading(doc, "8.2 销售预计", 2)
    add_table(
        doc,
        ["项目", "第一年", "第二年", "第三年"],
        [
            ("30天验证包数量", "3个", "8个", "15个"),
            ("标准点位数量", "1个", "4个", "8个"),
            ("年度服务续费", "0台", "1台", "5台"),
            ("转化抵扣", "1次×1500元", "3次×1500元", "6次×1500元"),
            ("预计营业收入", "23240元", "83740元", "168100元"),
        ],
        widths=[Cm(5.0), Cm(3.9), Cm(3.9), Cm(3.9)],
        font_size=8.8,
    )
    add_para(
        doc,
        "收入计算采用统一口径：验证包2980元；标准点位15800元；90天内由验证转采购时抵扣1500元；第二年起软件服务1200元每台。预测不包含医院信息系统、电梯控制和定制硬件收入。",
    )

    add_heading(doc, "8.3 成本费用核算", 2, new_page=True)
    add_table(
        doc,
        ["指标", "第一年", "第二年", "第三年"],
        [
            ("营业收入", "23240元", "83740元", "168100元"),
            ("直接交付成本", "13400元", "46700元", "92500元"),
            ("贡献毛额", "9840元", "37040元", "75600元"),
            ("毛额率", "42.3%", "44.2%", "45.0%"),
            ("研发销售管理费用", "10000元", "24000元", "42000元"),
            ("税前经营结果", "-160元", "13040元", "33600元"),
        ],
        widths=[Cm(5.0), Cm(3.9), Cm(3.9), Cm(3.9)],
        font_size=8.8,
    )
    add_para(
        doc,
        "直接成本按验证包1800元、标准点位8000元、年度服务300元估算；研发、销售和管理费用包含样机迭代、差旅沟通、宣传材料和工具服务。销售量、转化率或直接成本任一项偏离，经营结果都会改变，必须用真实试点数据逐季重算。",
    )

    add_heading(doc, "8.4 利润及利润分配", 2)
    add_para(
        doc,
        "项目未设立公司前不进行利润分配。若后续依法设立经营主体且实现累计可分配利润，前两个完整经营年度原则上不分红；现金优先用于产品安全、售后准备和下一轮试点。自第三个盈利年度起，在保留不少于六个月固定支出和全部保修准备金后，年度可分配利润拟按70%研发与质量改进、20%营运资金、10%团队激励配置，最终以公司章程、税务和成员书面决议为准。",
    )

    add_heading(doc, "8.5 融资方案和回报", 2)
    add_heading(doc, "8.5.1 资本结构与规模", 3)
    add_para(doc, "当前资本投入仅为不超过500元的参赛样机支出，未发行股权，也不存在投资人持股。计划书不编造注册资本、估值或股份比例。")
    add_heading(doc, "8.5.2 融资形式", 3)
    add_para(doc, "在取得至少1个付费验证、1份继续试用或采购意向且实车安全门槛通过前，不启动股权融资。优先申请校级创新资金、赛事奖金或企业联合验证经费，减少过早股权安排。")
    add_heading(doc, "8.5.3 资金来源与使用计划", 3)
    add_para(doc, "首轮外部非股权资金目标上限为5万元：40%用于结构和电子可靠性，30%用于安全测试、合规咨询与必要认证准备，20%用于受控试点和差旅培训，10%作为故障与售后准备金。所有支出按凭证和里程碑审核。")
    add_heading(doc, "8.5.4 投资收益与风险分析", 3, new_page=True)
    add_para(doc, "当前没有足够的真实销售与复购数据，不能承诺投资收益率或回收期。是否引入投资以四项门槛决定：完成付费验证、标准点位直接成本不超过8000元、30天试用无重大安全问题、机构明确继续采购或扩点。未满足时继续以科研和竞赛原型推进，不进入商业融资。")

    # 第九章
    add_heading(doc, "第九章 实施计划与阶段结论", 1, new_page=True)
    add_heading(doc, "9.1 硬件到货后的七天计划", 2)
    add_table(
        doc,
        ["日期", "重点任务", "当日完成条件"],
        [
            ("第1天", "清点器件、核对电压与引脚、固定模块", "形成实物接线表，上电无异常"),
            ("第2天", "单测电机、巡线、超声波和按键", "每个模块能独立读取或控制"),
            ("第3天", "标定黑白阈值和低速巡线参数", "直线和缓弯连续成功5次"),
            ("第4天", "实现障碍、丢线、超时和本地STOP", "五类安全测试全部停车"),
            ("第5天", "接入USB串口并测试蓝牙SPP", "网页命令与状态回传闭环"),
            ("第6天", "搭建医院路线并联调界面和小车", "完整流程连续成功10次"),
            ("第7天", "冻结功能、完成20次彩排和备用视频", "至少18次成功，演示少于2分钟"),
        ],
        widths=[Cm(2.2), Cm(7.5), Cm(6.9)],
        font_size=8.6,
    )

    add_heading(doc, "9.2 验证指标", 2)
    add_table(
        doc,
        ["指标", "当前状态", "比赛目标", "证据"],
        [
            ("自动测试", "28项通过", "保持全部通过", "pytest输出和提交号"),
            ("预设问法", "软件已覆盖", "15条全部正确", "输入输出记录"),
            ("跨楼层显示", "已完成", "去程和返程电梯提示正确", "路线步骤和楼层截图"),
            ("障碍停车", "待实车", "5次全部停车", "距离日志和视频"),
            ("丢线、断连、超时", "协议已设计", "逐项触发本地停车", "串口日志和视频"),
            ("完整演示", "软件闭环完成", "20次至少18次成功", "彩排表"),
            ("故障降级", "模拟模式完成", "30秒内切换", "计时记录"),
        ],
        widths=[Cm(3.2), Cm(3.3), Cm(5.0), Cm(5.1)],
        font_size=8.3,
    )

    add_heading(doc, "9.3 现场演示与降级预案", 2)
    add_table(
        doc,
        ["时间", "演示动作", "异常时的处理"],
        [
            ("0—20秒", "说明老年人多站就医的信息负担", "不扩展到诊断或陌生环境导航"),
            ("20—45秒", "输入先去心内科再去卫生间", "语音失败立即用按钮或文字"),
            ("45—80秒", "展示一楼到三楼及返程电梯步骤", "实体车只到电梯口，屏幕继续模拟"),
            ("80—105秒", "实体路线中放置软障碍并恢复", "状态异常立即STOP"),
            ("105—120秒", "到达并说明产品边界", "小车故障切换模拟，视频标注预录"),
        ],
        widths=[Cm(2.5), Cm(7.3), Cm(6.8)],
        font_size=8.5,
    )

    add_heading(doc, "9.4 风险矩阵", 2)
    add_table(
        doc,
        ["风险", "可能性", "影响", "预防与触发条件", "备用方案"],
        [
            ("实物型号不符", "中", "高", "到货先核对主控、电平与电源", "按实物改引脚，不强刷固件"),
            ("弯道丢线", "中", "高", "低速、宽弯、单变量标定", "缩短实体路线"),
            ("语音噪声", "高", "中", "显示识别文本并预演麦克风", "按钮和文字"),
            ("通信断开", "中", "高", "本地超时停车", "USB或模拟模式"),
            ("地图过期", "中", "高", "记录版本、审核人和更新责任", "停用旧路线"),
            ("能力被夸大", "中", "高", "材料区分已完成、待验证和规划", "删除无证据表述"),
        ],
        widths=[Cm(2.8), Cm(1.8), Cm(1.8), Cm(6.1), Cm(4.1)],
        font_size=7.9,
    )

    add_heading(doc, "9.5 阶段结论", 2)
    add_para(
        doc,
        "截至2026年9月10日，银伴软件演示版本已经完成，28项自动测试全部通过。五类地点问询、多目的地顺序解析、道路节点组合、跨楼层电梯过渡、分楼层路线动画、模拟机器人状态以及启动停止协议均有可复核结果。因此，软件端已经达到参赛演示验收线，可以独立完成两分钟核心流程。",
    )
    add_para(
        doc,
        "实体机器人尚未完成到货后的型号核验、接线、巡线标定、障碍停车和整机连续运行测试。因此，项目当前没有通过实体整机验收，不能宣称已经具备真实医院自主导引、自动乘梯或信息系统接入能力。当前产品性质确定为医院固定环境参赛原型。",
    )
    add_para(
        doc,
        "比赛展示方式按统一门槛决定：指定路线20次至少成功18次；5次障碍测试全部停车；人工STOP、丢线、通信中断和运行超时均能触发本地停机；完整演示不超过120秒；故障后30秒内能切换模拟模式。五项全部满足时采用软件与实物闭环演示；任一未满足时，实体车退出主流程，正式演示改用软件模拟，小车只作静态展示，备用视频明确标注为预录证据。",
    )

    # 附录
    add_heading(doc, "附录一 最小功能清单", 1, new_page=True)
    add_table(
        doc,
        ["功能", "状态", "验收方法"],
        [
            ("大按钮、文字、语音输入", "软件已完成", "三种方式提交同一目的地"),
            ("五类地点问询", "软件已完成", "心内科、药房、卫生间、检验科、挂号处均返回位置"),
            ("多目的地顺序", "软件已完成", "路线顺序与用户表达一致"),
            ("分楼层电梯步骤", "软件已完成", "一楼、三楼和返程转换明确"),
            ("屏幕模拟导引", "软件已完成", "图标沿节点移动并自动切换楼层"),
            ("单条实体巡线", "待硬件联调", "20次至少18次成功"),
            ("障碍与异常停车", "待硬件联调", "障碍、丢线、断连、超时和STOP逐项验证"),
            ("人工求助", "本地提醒已完成", "不得宣称已通知真实医护"),
        ],
        widths=[Cm(5.2), Cm(3.4), Cm(8.0)],
        font_size=8.6,
    )

    add_heading(doc, "附录二 软件目录", 1, new_page=True)
    directory = (
        "yinban-mvp\n"
        "  app\n"
        "    api       对话、导航与机器人接口\n"
        "    core      意图、路线、对话与安全规则\n"
        "    data      医院知识、楼层节点与路线数据\n"
        "    services  模拟机器人、串口与协议\n"
        "    web       适老界面、分楼层地图和前端逻辑\n"
        "  firmware\n"
        "    yinban_car  ESP32 Arduino固件\n"
        "  tests       自动测试\n"
        "  docs        接线、采购、演示和计划书\n"
        "  run.ps1     PowerShell启动脚本\n"
        "  start-yinban.cmd  Windows双击启动器"
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(1.2)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(directory), size=10, family="Consolas")

    add_heading(doc, "附录三 品牌文件", 1, new_page=True)
    logo = ASSET_DIR / "yinban-logo-horizontal-transparent.png"
    mascot = ASSET_DIR / "yinban-mascot-cover-transparent.png"
    if logo.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(logo), width=Cm(11.5))
        shape._inline.docPr.set("descr", "银伴项目横向标志，由亲和型机器人图形与银伴文字组成")
        add_caption(doc, "图7 银伴横向标志")
    if mascot.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(mascot), height=Cm(9.5))
        shape._inline.docPr.set("descr", "银伴项目亲和型机器人形象，采用圆润外观与暖黄色帽檐")
        add_caption(doc, "图8 银伴亲和型机器人形象")
    add_para(doc, "品牌使用墨绿、暖金、奶白和银灰作为主色。圆润机器人、路线环和定位标记共同表达陪伴与导引。图形目前用于参赛和项目展示，商标注册状态以4.1.4节为准。")

    add_heading(doc, "附录四 资料来源", 1, new_page=True)
    add_source(doc, 1, "中华人民共和国2025年国民经济和社会发展统计公报", "国家统计局", "2026年2月28日", "https://www.stats.gov.cn/xxgk/sjfb/tjgb2020/202602/t20260228_1962662.html")
    add_source(doc, 2, "关于切实解决老年人运用智能技术困难的实施方案", "国务院办公厅", "2020年", "https://www.nhc.gov.cn/bgt/gwywj2/202011/b51828e5adac4dbc92fbaa8d26474802.shtml")
    add_source(doc, 3, "改善就医感受提升患者体验主题活动方案（2023—2025年）", "国家卫生健康委、国家中医药局", "2023年", "https://www.nhc.gov.cn/wjw/c100375/202305/bfa59db84f1041c4b50bb859a3b76f39.shtml")
    add_source(doc, 4, "中华人民共和国无障碍环境建设法", "全国人民代表大会常务委员会", "2023年", "https://gxca.miit.gov.cn/zwgk/zcwj/flfg/art/2024/art_340a68cb6a8e45ae81c121630dd37346.html")
    add_source(doc, 5, "对十四届全国人大三次会议第7409号建议的答复", "国家卫生健康委", "2025年8月", "https://www.nhc.gov.cn/wjw/jiany/202508/c66483762ee2451fa550847bf28e0e89.shtml")
    add_source(doc, 6, "银伴项目代码仓库", "银伴项目团队", "2026年", "https://github.com/zhaiqianni/yinban-mvp")

    add_heading(doc, "附录五 答辩证据清单", 1, new_page=True)
    add_table(
        doc,
        ["检查项", "答辩前完成要求", "状态"],
        [
            ("软件版本", "冻结可运行版本并记录Git标签和提交号", "□"),
            ("自动测试", "保存完整测试输出，核心测试全部通过", "□"),
            ("跨楼层路线", "保留电梯去程、返程和地图切换截图", "□"),
            ("硬件核验", "保存实物型号、电压、引脚和接线照片", "□"),
            ("安全测试", "急停、障碍、丢线、超时和断连逐项留档", "□"),
            ("完整彩排", "20次至少18次成功并记录失败原因", "□"),
            ("备用方案", "USB、模拟模式和备用视频30秒内可切换", "□"),
            ("用户证据", "匿名访谈保留原始记录和反对意见", "□"),
            ("来源审查", "数字、政策、图片和代码均有来源", "□"),
            ("边界表述", "不宣称诊断、真实部署、自动乘梯或未验证能力", "□"),
        ],
        widths=[Cm(3.2), Cm(11.4), Cm(2.0)],
        font_size=8.7,
    )

    add_heading(doc, "附录六 用户访谈提纲", 1, new_page=True)
    add_table(
        doc,
        ["序号", "访谈问题", "记录重点"],
        [
            ("1", "请回忆最近一次到医院办事的完整过程。", "场景、时间顺序、是否有人陪同"),
            ("2", "其中最困难或最容易走错的是哪一步？", "地点、流程、理解与操作障碍"),
            ("3", "遇到困难时，通常向谁求助或用什么办法？", "人工、家属、手机、标识或放弃"),
            ("4", "一次要去多个地点时，怎样记住顺序？", "纸条、截图、口头提醒和改路线"),
            ("5", "语音、大按钮和文字更倾向哪一种？", "噪声、视力、识字和隐私"),
            ("6", "什么情况会让您愿意或拒绝跟随机器人？", "速度、距离、外观、人员和安全"),
            ("7", "机器人停止、迷路或无法回答时应怎样处理？", "人工接管、求助入口和等待"),
        ],
        widths=[Cm(1.7), Cm(8.1), Cm(6.8)],
        font_size=8.5,
    )
    add_para(doc, "访谈开始前说明用途、匿名方式和自愿退出权，不记录姓名、身份证、病历、诊断或联系方式。拒绝意见和失败样本与支持意见同样保留。")

    add_heading(doc, "附录七 两分钟演示操作清单", 1, new_page=True)
    add_table(
        doc,
        ["时间", "演示动作", "检查与异常分支"],
        [
            ("演示前", "启动服务、连接小车、清空旧任务", "检查电量、路线、音量、STOP和USB"),
            ("0—20秒", "说明数字化变化与老年就医信息负担", "不扩展到诊断和陌生环境导航"),
            ("20—45秒", "说出先去心内科再去卫生间", "语音失败改用按钮或文字"),
            ("45—80秒", "展示电梯上下楼和分层地图", "实体车到电梯口停止，屏幕续引"),
            ("80—105秒", "展示障碍停车和恢复", "状态异常立即STOP"),
            ("105—120秒", "到达并说明阶段边界", "故障改用模拟，视频注明预录"),
            ("结束后", "停止电机、保存日志、记录失败原因", "不在现场临时增加功能"),
        ],
        widths=[Cm(2.5), Cm(7.3), Cm(6.8)],
        font_size=8.5,
    )

    original_heading_signature = heading_signature(doc)
    append_section_enrichments(doc)
    if heading_signature(doc) != original_heading_signature:
        raise RuntimeError("内容扩充意外改变了原有标题结构或目录顺序。")

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    for section in doc.sections:
        configure_page(section)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
